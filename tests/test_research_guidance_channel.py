"""Owner, 2026-10-03: (i) qwen3.8-flash is the report engine, (ii) the thinking goes on
translating management guidance into financial estimates, (iii) the estimates reach the
report. Pins the model policy, the search strategy, the 2G section, the extractor's
normalisation and the DCF's guidance channel.
"""
import inspect
import re
from types import SimpleNamespace

import pytest

from src.agents.analysis import dcf_agent as d
from src.agents.industry import deep_research as dr
from src.agents.industry import sector_prompts as sp
from src.memory import run_archive as ra


# ── (i) the model and its thinking policy ───────────────────────────────────

def test_the_report_engine_is_the_flash_tier_and_the_routing_defaults_to_it():
    assert dr.DEFAULT_RESEARCH_MODEL == "qwen3.8-flash"
    src = inspect.getsource(dr.run_deep_research_agent)
    assert 'or "qwen3.6-plus"' not in src
    assert src.count("DEFAULT_RESEARCH_MODEL") >= 4          # US default, HK search, HK synthesis, US-on-Qwen synthesis


def test_flash_calls_get_thinking_off_unless_a_budget_is_asked_for():
    kw = {"model": "qwen3.8-flash", "max_tokens": 400, "messages": []}
    out = dr._apply_thinking_policy(kw)
    assert out["thinking"] == {"type": "disabled"} and "thinking" not in kw   # a copy, the caller's dict untouched
    # an asked-for budget is kept, and max_tokens is lifted above it so the answer is not cut
    kw2 = {"model": "qwen3.8-flash", "max_tokens": 400, "thinking": {"type": "enabled", "budget_tokens": 6000}}
    out2 = dr._apply_thinking_policy(kw2)
    assert out2["thinking"]["budget_tokens"] == 6000 and out2["max_tokens"] == 8500
    # other models are untouched
    for m in ("qwen3.6-plus", "claude-sonnet-4-6", "claude-opus-5", None):
        assert dr._apply_thinking_policy({"model": m, "max_tokens": 400}) == {"model": m, "max_tokens": 400}
    assert dr._apply_thinking_policy(out) == out                              # idempotent


def test_the_estimates_pass_thinks_with_a_budget_on_qwen_and_adaptively_on_claude(monkeypatch):
    monkeypatch.delenv("GUIDANCE_THINKING_BUDGET", raising=False)
    assert dr._thinking_for_estimates("qwen3.8-flash") == {"type": "enabled", "budget_tokens": 6000}
    assert dr._thinking_for_estimates("claude-sonnet-4-6") == {"type": "adaptive"}
    monkeypatch.setenv("GUIDANCE_THINKING_BUDGET", "3000")
    assert dr._thinking_for_estimates("qwen3.6-plus")["budget_tokens"] == 3000


def test_every_client_deep_research_builds_carries_the_policy():
    calls = []

    class _Msgs:
        def create(self, **kw):
            calls.append(kw)
            return SimpleNamespace(content=[])

    class _Client:
        def __init__(self):
            self.messages = _Msgs()

    c = dr._with_thinking_policy(_Client())
    c.messages.create(model="qwen3.8-flash", max_tokens=100)
    assert calls[-1]["thinking"] == {"type": "disabled"}
    c.messages.create(model="qwen3.6-plus", max_tokens=100)
    assert "thinking" not in calls[-1]
    # and the retry wrapper applies it for raw clients too
    dr._call_llm_with_rate_retry(_Client(), extractor_name="x", model="qwen3.8-flash", max_tokens=50)
    assert calls[-1]["thinking"] == {"type": "disabled"}
    src = inspect.getsource(dr)
    assert src.count("= anthropic.Anthropic(") == 1                           # only inside make_sdk_client
    assert "make_sdk_client(" in inspect.getsource(dr._research_one_ticker)


def test_the_search_strategy_is_per_model(monkeypatch):
    monkeypatch.delenv("DEEP_RESEARCH_SEARCH_STRATEGY", raising=False)
    monkeypatch.delenv("DEEP_RESEARCH_SEARCH_THINKING", raising=False)
    assert dr._qwen_search_strategy("qwen3.8-flash") == "pro"                # "agent" is refused by the flash tier
    assert dr._qwen_search_strategy("qwen3.6-plus") == "agent"
    body = dr._qwen_search_extra_body("qwen3.8-flash")
    assert body == {"enable_search": True, "search_options": {"search_strategy": "pro"}, "enable_thinking": False}
    assert dr._qwen_search_extra_body("qwen3.6-plus") == {"enable_search": True, "search_options": {"search_strategy": "agent"}}
    monkeypatch.setenv("DEEP_RESEARCH_SEARCH_STRATEGY", "turbo")
    monkeypatch.setenv("DEEP_RESEARCH_SEARCH_THINKING", "on")
    body = dr._qwen_search_extra_body("qwen3.8-flash")
    assert body["search_options"] == {"search_strategy": "turbo"} and "enable_thinking" not in body
    assert '"search_strategy": "agent"' not in inspect.getsource(dr._research_one_ticker)


# ── (ii) the research asks for guidance first and writes 2G ─────────────────

def test_the_system_prompt_leads_with_guidance_and_specifies_2g(monkeypatch):
    monkeypatch.delenv("DEEP_RESEARCH_SEARCH_PROFILE", raising=False)
    p = dr._build_research_system("2026", "Tech", "Mature SaaS")
    assert "2G. MANAGEMENT GUIDANCE" in p and "GUIDANCE_BLOCK_START" in p and "GUIDANCE_BLOCK_END" in p
    assert p.index("GUIDANCE FIRST") < p.index('"[Ticker] market share competitive landscape')
    assert "2A through 2G" in p and "2A through 2F" not in p
    assert "2G (Guidance)" in p
    assert "bear <= base <= bull" in p.lower().replace("≤", "<=") or "BEAR / BASE / BULL" in p


def test_both_section_parsers_capture_2g():
    text = ("## 2A. Profit pool\nalpha\n\n## 2F. KPIs\nkpis\n\n## 2G. Management guidance → estimates\n"
            "GUIDANCE_BLOCK_START\nGUIDANCE | metric=revenue | period=FY2026 | low=290 | mid=292.5 | high=295 | unit=USD bn\n"
            "GUIDANCE_BLOCK_END\n\nSECTION 7 — INDUSTRY INTELLIGENCE BRIEF\nbrief")
    secs = dr._extract_sections(text)
    assert "2g" in secs and secs["2g"].startswith("## 2G.") and "brief" in secs
    assert "INDUSTRY INTELLIGENCE BRIEF" not in secs["2g"]
    inline = ra._parse_sections_inline(text)
    assert "2g" in inline
    assert sp.needs_extractor("guidance_estimates", "RealEstate", "R.E.I.T.") is True
    assert sp.needs_extractor("guidance_estimates", "", "") is True


# ── the extractor's normalisation ───────────────────────────────────────────

_PARSED = {
    "as_of": "2026-09-25", "fiscal_year_1": "FY2026", "fiscal_year_2": "FY2027",
    "guidance": {"revenue_growth": {"low": "5.4%", "mid": 0.063, "high": 7.2}, "revenue": {"low": 290, "mid": 292.5, "high": 295, "currency": "USD", "scale": "bn"},
                 "ebitda_margin": {"low": 0.112, "mid": 0.1145, "high": 0.117}, "eps": None, "basis": "reported", "status": "raised",
                 "quote": "We now expect revenue of $290 to $295 billion", "source": "Q4 FY2025 release, 2026-09-25"},
    "consensus": {"revenue_growth_fy1": 0.058, "eps_fy1": 18.4, "as_of": "2026-09-30", "source": "FMP"},
    "guidance_vs_consensus_pct": 0.005,
    "track_record": "beat the midpoint 7 of the last 8 quarters",
    "estimates": {"bull": {"revenue_growth_fy1": 0.054, "revenue_growth_fy2": 0.05, "ebitda_margin_fy1": 0.112},   # mis-ordered on purpose
                  "base": {"revenue_growth_fy1": 0.065, "revenue_growth_fy2": 0.06, "ebitda_margin_fy1": 0.115, "eps_fy1": 18.9},
                  "bear": {"revenue_growth_fy1": 0.072, "revenue_growth_fy2": 0.065, "ebitda_margin_fy1": 0.118}},
    "rationale": "Base sits above the midpoint for the beat record; bull at the top of the range.",
    "confidence": "high", "citations": ["Q4 FY2025 release, 2026-09-25", "FMP consensus 2026-09-30"],
}


def test_the_extractor_normalises_rates_orders_scenarios_and_keeps_guidance_pure():
    out = dr._normalize_guidance_estimates(_PARSED)
    assert out["guidance"]["revenue_growth"] == {"low": 0.054, "mid": 0.063, "high": 0.072}     # "5.4%" and 7.2 read as percents
    assert out["guidance"]["revenue"]["scale"] == "bn" and out["guidance"]["eps"] is None
    est = out["estimates"]
    assert [est[s]["revenue_growth_fy1"] for s in ("bear", "base", "bull")] == [0.054, 0.065, 0.072]   # sorted bear <= base <= bull
    assert [est[s]["revenue_growth_fy2"] for s in ("bear", "base", "bull")] == [0.05, 0.06, 0.065]
    assert est["base"]["eps_fy1"] == 18.9 and est["bear"]["eps_fy1"] is None                           # an incomplete field is not sorted
    assert out["confidence"] == "HIGH" and out["consensus"]["revenue_growth_fy1"] == 0.058
    assert out["citations"] == _PARSED["citations"]
    # nothing to channel -> {}
    assert dr._normalize_guidance_estimates({"estimates": {"base": {"revenue_growth_fy1": None}}}) == {}
    assert dr._normalize_guidance_estimates({"estimates": {"base": {"revenue_growth_fy1": 9.0}}})["estimates"]["base"]["revenue_growth_fy1"] == 0.09   # percent-looking
    assert dr._normalize_guidance_estimates({"estimates": {"base": {"revenue_growth_fy1": 900.0}}}) == {}    # out of range is dropped
    assert dr._normalize_guidance_estimates("garbage") == {}


def test_the_extractor_reads_2g_thinks_and_skips_thinking_blocks():
    seen = {}

    class _Msgs:
        def create(self, **kw):
            seen.update(kw)
            return SimpleNamespace(content=[SimpleNamespace(type="thinking", thinking="..."),
                                            SimpleNamespace(type="text", text='{"estimates": {"base": {"revenue_growth_fy1": 0.06}}, "confidence": "MEDIUM"}')])

    client = SimpleNamespace(messages=_Msgs())
    out = dr._extract_guidance_estimates(client, "qwen3.8-flash", {"2g": "GUIDANCE_BLOCK_START ... GUIDANCE_BLOCK_END"}, "", "COST")
    assert out["estimates"]["base"]["revenue_growth_fy1"] == 0.06 and out["_model"] == "qwen3.8-flash"
    assert seen["thinking"]["type"] == "enabled" and seen["max_tokens"] > seen["thinking"]["budget_tokens"]
    assert "temperature" not in seen                                           # thinking and temperature do not mix
    assert "Section 2G" in seen["messages"][0]["content"]
    # a report with no guidance at all -> {} and no LLM call
    seen.clear()
    assert dr._extract_guidance_estimates(client, "qwen3.8-flash", {"2a": "profit pool"}, "nothing relevant here", "X") == {}
    assert not seen
    # legacy report (no 2G): the paragraphs about guidance are the source
    src = dr._guidance_source_text({"2a": "x"}, "Intro.\n\nManagement guided FY26 revenue up 6%.\n\nUnrelated.")
    assert "guided" in src and "Unrelated" not in src
    assert "guidance_estimates" in dr._RETRY_ON_EMPTY_EXTRACTORS


# ── the DCF's guidance channel ──────────────────────────────────────────────

_EST = {"confidence": "HIGH", "fiscal_year_1": "FY2026", "fiscal_year_2": "FY2027",
        "estimates": {"bear": {"revenue_growth_fy1": 0.03, "revenue_growth_fy2": 0.035},
                      "base": {"revenue_growth_fy1": 0.065, "revenue_growth_fy2": 0.06},
                      "bull": {"revenue_growth_fy1": 0.09, "revenue_growth_fy2": None}}}
_CFG = {"explicit_years": 2, "fade_years": 3, "min_confidence": "MEDIUM", "growth_cap": 0.40, "growth_floor": -0.30}


def test_the_channel_sets_the_explicit_years_and_fades_onto_the_engine_path(monkeypatch):
    monkeypatch.delenv("GUIDANCE_CHANNEL", raising=False)
    out = d._guidance_channel_schedule(_EST, "base", 0.10, None, 10, _CFG)
    s = out["schedule"]
    assert s[:2] == [0.065, 0.06]                                               # FY+1 then FY+2
    assert s[2] == pytest.approx(0.06 * 0.75 + 0.10 * 0.25) and s[3] == pytest.approx(0.06 * 0.5 + 0.10 * 0.5)
    assert s[4] == pytest.approx(0.06 * 0.25 + 0.10 * 0.75) and s[5:] == [0.10] * 5
    assert out["engine_year1"] == 0.10 and out["explicit"] == [0.065, 0.06]
    # a scenario without FY+2 repeats FY+1; the fade target is the engine's OWN schedule when it has one
    eng = [0.20, 0.18, 0.16, 0.14, 0.12, 0.10, 0.08, 0.06, 0.04, 0.03]
    bull = d._guidance_channel_schedule(_EST, "bull", 0.20, eng, 10, _CFG)
    assert bull["schedule"][:2] == [0.09, 0.09] and bull["schedule"][5:] == eng[5:]
    assert bull["schedule"][2] == pytest.approx(0.09 * 0.75 + 0.16 * 0.25)
    # bear takes the bear estimate
    assert d._guidance_channel_schedule(_EST, "bear", 0.10, None, 10, _CFG)["schedule"][0] == 0.03


def test_the_channel_declines_low_confidence_missing_rows_and_the_kill_switch(monkeypatch):
    monkeypatch.delenv("GUIDANCE_CHANNEL", raising=False)
    low = {**_EST, "confidence": "LOW"}
    assert d._guidance_channel_schedule(low, "base", 0.1, None, 10, _CFG) is None
    assert d._guidance_channel_schedule({}, "base", 0.1, None, 10, _CFG) is None
    assert d._guidance_channel_schedule(None, "base", 0.1, None, 10, _CFG) is None
    no_row = {"confidence": "HIGH", "estimates": {"base": {"revenue_growth_fy1": None}}}
    assert d._guidance_channel_schedule(no_row, "base", 0.1, None, 10, _CFG) is None
    # the clip: a 90% estimate is held at the cap
    hot = {"confidence": "HIGH", "estimates": {"base": {"revenue_growth_fy1": 0.9}}}
    assert d._guidance_channel_schedule(hot, "base", 0.1, None, 10, _CFG)["schedule"][0] == 0.40
    monkeypatch.setenv("GUIDANCE_CHANNEL", "off")
    assert d._guidance_channel_schedule(_EST, "base", 0.1, None, 10, _CFG) is None
    payload = d._guidance_estimates_payload(_EST, None)
    assert payload["applied"] is False and "GUIDANCE_CHANNEL is off" in payload["not_applied_reason"]


def test_the_constants_are_proposed_and_read_by_the_engine():
    from src.data import valuation_constants as vc
    blk = vc.load()["guidance_channel"]
    assert blk["status"] == "PROPOSED" and blk["reviewed"]["reviewer"] is None
    cfg = d._guidance_channel_cfg()
    assert (cfg["explicit_years"], cfg["fade_years"], cfg["min_confidence"]) == (2, 3, "MEDIUM")
    assert d._guidance_estimates_payload(None, None) is None
    p = d._guidance_estimates_payload({**_EST, "_model": "qwen3.8-flash"},
                                      d._guidance_channel_schedule(_EST, "base", 0.1, None, 10, cfg))
    assert p["applied"] is True and p["model"] == "qwen3.8-flash" and p["channel"]["explicit"] == [0.065, 0.06]
    assert "_model" not in p


def test_the_engine_wires_the_channel_into_the_scenario_loop_and_the_payload():
    src = inspect.getsource(d.run_dcf_agent)
    assert ("_guidance_channel_schedule(" + chr(10) + "                _guid_est, scenario, g, _growth_schedule, _PROJECTION_YEARS," + chr(10)
            + "                growth_adj=_guidance_growth_adj_for(ticker, sector))") in src
    assert '"guidance_estimates": _guidance_estimates_payload(_guid_est, _gc_applied),' in src
    assert '"guidance_channel": _gc,' in src
    assert 'state["data"].get("guidance_estimates", {})' in src
    # the research hands the block to the state under the same key
    rsrc = inspect.getsource(dr.run_deep_research_agent)
    assert 'state["data"]["guidance_estimates"] = guidance_estimates_all' in rsrc
    assert re.search(r'"guidance_estimates":\s+_ext_results\.get\("guidance_estimates", \{\}\)', inspect.getsource(dr._research_one_ticker))


# ── report length and the estimates in the PDF and the workbook ─────────────

def test_the_fast_tier_report_is_capped_and_the_compact_budget_is_the_default(monkeypatch):
    monkeypatch.delenv("DEEP_RESEARCH_REPORT_MAX_TOKENS", raising=False)
    monkeypatch.delenv("DEEP_RESEARCH_LENGTH", raising=False)
    monkeypatch.delenv("DEEP_RESEARCH_SEARCH_PROFILE", raising=False)
    monkeypatch.setenv("INDUSTRY_BRIEF_MODE", "merged")
    assert dr._report_max_tokens("qwen3.8-flash") == 12000 and dr._report_max_tokens("qwen3.6-plus") == dr.MAX_TOKENS
    monkeypatch.setenv("DEEP_RESEARCH_REPORT_MAX_TOKENS", "9000")
    assert dr._report_max_tokens("claude-sonnet-4-6") == 9000
    src = inspect.getsource(dr._research_one_ticker)
    assert src.count("max_tokens=_report_max_tokens(model_name)") == 5         # Tier-1 call, its nudge, the Qwen stream, its thinking-on retry, its synthesis retry
    p = dr._build_research_system("2026", "Tech", "")
    assert "LENGTH BUDGET (compact" in p and "8 to 12 bullets" in p          # owner: the brief keeps its depth
    assert "competitive position and share dynamics" in p and "the 2G estimates" in p
    assert "Footnote numbers are GLOBAL" in p
    monkeypatch.setenv("DEEP_RESEARCH_LENGTH", "full")
    p2 = dr._build_research_system("2026", "Tech", "")
    assert "LENGTH BUDGET" not in p2 and "8 to 12 bullets" in p2


_PAYLOAD = {"as_of": "2026-09-25", "fiscal_year_1": "FY2026", "fiscal_year_2": "FY2027", "confidence": "HIGH", "applied": True,
            "guidance": {"revenue_growth": {"low": 0.054, "mid": 0.063, "high": 0.072}, "revenue": None,
                         "ebitda_margin": {"low": 0.112, "mid": None, "high": 0.117}, "eps": None,
                         "basis": "reported", "status": "raised", "quote": "We now expect revenue of $290 to $295 billion", "source": "Q4 release"},
            "consensus": {"revenue_growth_fy1": 0.058, "eps_fy1": 18.4, "as_of": "2026-09-30"},
            "guidance_vs_consensus_pct": 0.005, "track_record": "beat 7 of 8",
            "estimates": {"bear": {"revenue_growth_fy1": 0.054, "revenue_growth_fy2": 0.05, "ebitda_margin_fy1": 0.112, "eps_fy1": None},
                          "base": {"revenue_growth_fy1": 0.065, "revenue_growth_fy2": 0.06, "ebitda_margin_fy1": 0.115, "eps_fy1": 18.9},
                          "bull": {"revenue_growth_fy1": 0.072, "revenue_growth_fy2": 0.065, "ebitda_margin_fy1": 0.118, "eps_fy1": 19.4}},
            "rationale": "Base above the midpoint for the beat record.", "citations": ["Q4 release", "FMP"],
            "channel": {"explicit": [0.065, 0.06], "explicit_years": 2, "fade_years": 3, "engine_year1": 0.1,
                        "schedule": [0.065, 0.06, 0.07, 0.08, 0.09, 0.1, 0.1, 0.1, 0.1, 0.1]}}


def test_the_pdf_prints_the_estimates_block_only_when_there_is_one():
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from src.utils import pdf_report as pr
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="RptSubsection", parent=styles["Heading3"]))
    assert pr._guidance_estimates_block_pdf({}, styles, 400.0) == []
    assert pr._guidance_estimates_block_pdf({"guidance_estimates": {"estimates": {}}}, styles, 400.0) == []
    out = pr._guidance_estimates_block_pdf({"guidance_estimates": _PAYLOAD}, styles, 400.0)
    assert len(out) >= 4
    table = [o for o in out if o.__class__.__name__ == "Table"][0]
    labels = [row[0].text for row in table._cellvalues[1:]]
    assert labels[0].startswith("Revenue growth FY2026") and any(l.startswith("EPS FY2026") for l in labels)
    assert not any(l.startswith("EPS FY2027") for l in labels)                 # no values anywhere -> row dropped
    texts = " ".join(getattr(o, "text", "") for o in out)
    assert "years 1–2" in texts and "+10.0%" in texts and "beat 7 of 8" in texts
    assert "_guidance_estimates_block_pdf(dcf_ticker, styles, page_w)" in inspect.getsource(pr)


def test_the_workbook_gets_a_guidance_tab_when_the_research_produced_estimates():
    import openpyxl
    from src.utils import valuation_workbook as vw
    book = vw._Book.__new__(vw._Book)
    book.wb, book.tabs, book.dr = openpyxl.Workbook(), [], {"guidance_estimates": _PAYLOAD}
    book.guidance_tab()
    ws = book.wb["Guidance"]
    col_a = [str(c.value) for c in ws["A"] if c.value is not None]
    assert any(v.startswith("Revenue growth FY2026") for v in col_a) and any(v.startswith("Consensus (FY2026)") for v in col_a)
    assert any("Base DCF growth path" in v for v in col_a)
    pct_cells = [c for row in ws.iter_rows() for c in row if c.number_format == vw.PCT and isinstance(c.value, float)]
    assert any(abs(c.value - 0.065) < 1e-9 for c in pct_cells)                  # the base FY+1 estimate, as a number
    assert ("Guidance", "Management guidance → the model's estimates and how the DCF used them") in book.tabs
    assert 'if (self.dr or {}).get("guidance_estimates"):' in inspect.getsource(vw._Book.build)
    assert 'self.guidance_tab()' in inspect.getsource(vw._Book.build)


# ── the native-search prompt and the degenerate Tier-1 guard (TSCO smoke, 2026-10-03) ──

def test_the_native_search_prompt_has_no_tool_to_call_and_the_tool_prompt_is_unchanged():
    native = dr._build_research_system("2026", "Consumer", "Traditional Retail", native_search=True)
    tool = dr._build_research_system("2026", "Consumer", "Traditional Retail")
    assert "NO tool to call" in native and "[TOOL CALL]" in native and "exactly ONE tool" not in native
    assert "MUST call web_search" not in native and "2G. MANAGEMENT GUIDANCE" in native
    assert "exactly ONE tool: web_search" in tool and "MUST call web_search at least 8 times" in tool
    assert "[B] WEB SEARCH (your tool)" in tool and "[B] WEB SEARCH (native, already run)" in native
    src = inspect.getsource(dr._research_one_ticker)
    assert src.count('{"role": "system", "content": _research_system_native}') == 2      # first try and the thinking-on retry
    assert src.count('{"role": "system",    "content": _research_system_native}') == 1  # the non-streaming synthesis retry
    assert '"enable_thinking") is False' in src and "_consume_stream(_stream)" in src


def test_a_textual_tool_call_or_a_stub_is_degenerate_and_a_report_is_not():
    assert dr._tier1_degenerate('[TOOL CALL] {"name": "web_search", "arguments": {"query": "TSCO guidance"}}')
    assert dr._tier1_degenerate("")
    nl = chr(10)
    assert dr._tier1_degenerate("2A. Profit pool" + nl + "x " * 400)                        # 800 chars: too short
    report = "## 2A. Profit pool" + nl + ("Revenue grew 13% to $69.7B [1]. " * 80) + nl + "## 2G. Guidance" + nl + "y " * 200
    assert not dr._tier1_degenerate(report)
    assert dr._tier1_degenerate('{"name": "web_search", "arguments": {}}' + nl + "z " * 1000)   # long, but a tool call up front


# ── the brief carries its citations (owner, 2026-10-03: the research text is not printed) ──

_SECTIONS = {
    "2a": "## 2A. Profit pool" + chr(10) + "Revenue grew 13% [1]. Share 40% [2]." + chr(10) + "REFERENCES" + chr(10)
          + "[1] Accenture Form 10-K FY2025, SEC EDGAR, acc: 0001467373-25-000082" + chr(10)
          + "    URL: https://www.sec.gov/Archives/edgar/x" + chr(10)
          + "[2] IoT Analytics — GenAI Services Market Share Report, January 2025" + chr(10)
          + "    URL: URL unavailable — paywalled",
    "2g": "## 2G. Guidance" + chr(10) + "FY26 revenue $290-295bn [3]." + chr(10) + "REFERENCES" + chr(10)
          + "[3] Jensen Huang, CEO — Q3 FY2026 Earnings Call, 20 November 2025" + chr(10) + "    URL: https://investor.x/call",
    "brief": "SECTION 7" + chr(10) + "- Guidance above the Street [3]; share gains [2] on the 10-K base [1]; unsourced claim [9].",
}


def test_the_brief_markers_resolve_to_the_sections_references_in_order():
    entries = dr.parse_reference_entries(_SECTIONS)
    assert sorted(entries) == [1, 2, 3]
    assert entries[1]["url"] == "https://www.sec.gov/Archives/edgar/x" and entries[1]["verified"] is True
    assert entries[2]["url"] is None and entries[2]["date"] == "January 2025"          # paywalled: no url, the date read
    assert entries[3]["date"] == "20 November 2025" and entries[3]["section"] == "2g"
    rows = dr.resolve_brief_references(_SECTIONS["brief"], _SECTIONS)
    assert [r["ref_id"] for r in rows] == [3, 2, 1, 9]                                   # order of first appearance
    assert rows[0]["source_name"].startswith("Jensen Huang") and rows[3]["source_type"] == "unresolved"
    assert dr.resolve_brief_references("", _SECTIONS) == [] and dr.resolve_brief_references("no markers", {}) == []
    # a legacy report that restarted numbering: the first entry with a number wins
    dup = {"2a": "x" + chr(10) + "REFERENCES" + chr(10) + "[1] First source, 2025", "2b": "y" + chr(10) + "REFERENCES" + chr(10) + "[1] Second source, 2026"}
    assert dr.parse_reference_entries(dup)[1]["source_name"] == "First source, 2025"
    from src.agents.industry import specialist as sp_
    src = inspect.getsource(sp_.assemble_industry_brief_merged)
    assert "resolve_brief_references" in src and 'state["data"]["industry_footnotes"] = []' not in src.split("except")[0]
    # a brief loaded from the archive's phase cache resolves its footnotes too (prod 2026-10-03:
    # four cached runs carried a brief with [n] markers and no footnotes)
    from pathlib import Path
    pipe = (Path(__file__).resolve().parents[1] / "src" / "pipeline.py").read_text(encoding="utf-8")
    assert pipe.index("Industry brief loaded from archive") > pipe.index("resolve_brief_references as _rbr")


def test_the_est_calibration_shifts_the_explicit_years_and_zero_is_bit_identical(monkeypatch):
    """Self-learning loop 1: the ACTIVE est calibration's adjustment reaches the channel as
    growth_adj; 0.0 (nothing promoted) reproduces the schedule exactly."""
    monkeypatch.delenv("GUIDANCE_CHANNEL", raising=False)
    base = d._guidance_channel_schedule(_EST, "base", 0.10, None, 10, _CFG)
    same = d._guidance_channel_schedule(_EST, "base", 0.10, None, 10, _CFG, growth_adj=0.0)
    assert same == base and "growth_adj" not in same
    low = d._guidance_channel_schedule(_EST, "base", 0.10, None, 10, _CFG, growth_adj=-0.02)
    assert low["explicit"] == [pytest.approx(0.045), pytest.approx(0.04)]
    assert low["schedule"][0] == pytest.approx(0.045) and low["growth_adj"] == -0.02
    assert low["schedule"][5:] == base["schedule"][5:]                      # the engine path is untouched
    # the reader is inert until an est version is active, and never raises
    from src.memory import calibration as cal
    monkeypatch.setattr(cal, "active_version", lambda family="iv": None)
    assert d._guidance_growth_adj_for("ZZCO", "Tech") == 0.0
    monkeypatch.setattr(cal, "active_version", lambda family="iv": {"version_id": "calest-x", "family": "est",
                        "params": {"guidance_growth_adj": {"sector:Tech": -0.02, "market:US": 0.01}}})
    assert d._guidance_growth_adj_for("ZZCO", "Tech") == -0.02               # sector scope wins
    assert d._guidance_growth_adj_for("ZZCO", "Energy") == 0.01              # then the market
    assert d._guidance_growth_adj_for("0005.HK", None) == 0.0
