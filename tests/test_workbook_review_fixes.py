"""SBUX workbook review (owner, 2026-10-03), the quick fixes: no empty SOTP tab (A10), the
probability-weighted and base-case targets each named (A9), every net-debt figure with its basis
(A6), an unweighted DCF that says so and reconciles (A1), the leverage test marked n/a on
negative equity and the WACC the valuation used stated with its parts (A8)."""
import importlib.util
import io
from pathlib import Path

from openpyxl import load_workbook

from src.utils.valuation_workbook import build_workbook

_spec = importlib.util.spec_from_file_location("_wbtests", Path(__file__).resolve().parent / "test_valuation_workbook.py")
_wbt = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(_wbt)


def _col_a(ws):
    return [str(c.value) for c in ws["A"] if c.value is not None]


def _run_without_dcf_weight():
    """A multiples-only blend (the Restaurants shape) with an unweighted SOTP (segments) trace and
    negative book equity in the WACC build."""
    sotp = {"kind": "sotp", "source": "segments", "value": 71.16, "table": {"rows": [{"name": "x", "method": "EV/EBITDA", "multiple": 10.0, "value": 1000.0}],
                                                                             "associates": 0.0, "net_cash": -100.0, "nav": 900.0, "shares": 10.0}}
    wts = [{"method": "EV/EBITDA", "value_key": "EV/EBITDA", "bucket": "multi", "weight": 0.6},
           {"method": "P/E", "value_key": "P/E", "bucket": "multi", "weight": 0.4}]
    run = _wbt._run_with({"SOTP (segments)": sotp}, wts)
    dr = run["data"]["dcf_range"]["TEST"]
    dr["wacc_build"]["base_breakdown"]["leverage"] = -3.29
    dr["wacc_build"]["base_breakdown"]["leverage_premium_applies"] = True
    dr["wacc_build"]["base_breakdown"]["macro_overlay"] = 0.0025
    dr["wacc_build"]["base_breakdown"]["macro_regime"] = "risk-off"
    dr["wacc"] = 0.0925
    dr["financials_used"] = {"balance_sheet_period": "2026-06-28"}
    return run


def test_an_unweighted_sotp_trace_creates_no_tab_and_an_unweighted_dcf_says_so():
    wb = load_workbook(io.BytesIO(build_workbook(_run_without_dcf_weight(), "TEST", load_statements=_wbt._statements)))
    assert "SOTP" not in wb.sheetnames
    assert "SOTP" not in _col_a(wb["Cover"])
    # Owner, 2026-10-04 (decision D1): an unweighted DCF is not an output -- no tab, one Summary line.
    assert "DCF" not in wb.sheetnames
    notes = [v for v in _col_a(wb["Summary"]) if v.startswith("DCF EXCLUDED from the blend")]
    assert len(notes) == 1 and "Test Profile profile weights EV/EBITDA, P/E" in notes[0]
    assert "blended intrinsic value" in notes[0]
    # the default fixture weights the DCF: the tab stands and no such line
    wb2 = load_workbook(io.BytesIO(build_workbook(_wbt._run(), "TEST")))
    assert "DCF" in wb2.sheetnames
    assert not [v for v in _col_a(wb2["Summary"]) if v.startswith("DCF EXCLUDED")]


def test_the_two_targets_and_the_three_net_debts_are_each_named():
    wb = load_workbook(io.BytesIO(build_workbook(_run_without_dcf_weight(), "TEST", load_statements=_wbt._statements)))
    s = wb["Summary"]
    labels = _col_a(s)
    assert "12-month target (probability-weighted)" in labels and "Base-case target (spot + capture × (IV − spot))" in labels
    r_pw = next(c.row for c in s["A"] if c.value == "12-month target (probability-weighted)")
    r_base = next(c.row for c in s["A"] if c.value == "Base-case target (spot + capture × (IV − spot))")
    assert str(s.cell(row=r_pw, column=3).value).startswith("='Target'!$C$9")
    assert s.cell(row=r_base, column=3).value == "='Target'!$C$6"
    from src.data.report_families import _R
    assert _R["net_debt"][0] == "Net debt / (cash) — FMP annual, lease liabilities included"   # the family history row's label
    # Alibaba review (2026-10-03, section 1): the bridge deducts minorities and preferreds on both tabs,
    # and the Summary's static column carries the same sign as the linked "Less:" cells.
    assert "DCF" not in wb.sheetnames                 # D1: this fixture's DCF carries no weight
    _r3 = _wbt._run()
    _r3["data"]["dcf_range"]["TEST"]["financials_used"] = {"balance_sheet_period": "2026-06-28", "net_debt_basis": {
        "balance_sheet_date": "2026-06-28", "short_term_investments_netted": True, "accounting_basis": "US GAAP",
        "leases": "excluded (US GAAP: rent is inside EBITDA and cash flow)"}}
    wb3 = load_workbook(io.BytesIO(build_workbook(_r3, "TEST", load_statements=_wbt._statements)))
    s3 = wb3["Summary"]
    assert "Less: minority interest" in _col_a(s3) and "Less: preferred equity" in _col_a(s3)
    # Owner, 2026-10-04 (plan 1A.3): the label states the recorded basis.
    assert any(v.startswith("Less: net debt — valuation basis (balance sheet 2026-06-28; short-term investments counted as cash; "
                            "US GAAP: leases excluded") for v in _col_a(s3))
    assert "Less: minority interest" in _col_a(wb3["DCF"]) and "Less: preferred equity" in _col_a(wb3["DCF"])
    r_nd = next(c.row for c in s3["A"] if str(c.value).startswith("Less: net debt — valuation basis"))
    _static, _linked = s3.cell(row=r_nd, column=2).value, s3.cell(row=r_nd, column=3).value
    _nd_trace = _wbt._run()["data"]["dcf_range"]["TEST"]["base"]["leg_inputs"]["DCF"]["net_debt"]
    assert _static == -_nd_trace                      # "Less:" sign, matching the DCF tab's =-nd cell
    assert str(_linked).startswith("='DCF'!") or str(_linked).startswith("=DCF!")
    assert any(v.startswith("Net debt — this tab's formula (debt − cash − short-term investments; leases excluded)") for v in _col_a(wb["BS"]))


def test_negative_equity_marks_the_leverage_test_na_and_the_backtest_states_the_wacc_used():
    wb = load_workbook(io.BytesIO(build_workbook(_run_without_dcf_weight(), "TEST", load_statements=_wbt._statements)))
    w = wb["WACC"]
    row = next(c for c in w["A"] if isinstance(c.value, str) and c.value.strip().startswith("Leverage premium applies"))
    assert "n/a: book equity is negative" in row.value and w.cell(row=row.row, column=3).value == 0
    bt = _col_a(wb["Backtest"])
    line = next(v for v in bt if v.startswith("WACC used by the valuation"))
    assert "9.25%" in line and "build 9.00%" in line and "macro-regime overlay +0.25%" in line and "risk-off regime" in line
