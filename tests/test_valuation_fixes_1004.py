"""Plan 2026-10 Phase 1 (owner, 2026-10-04): the engine-wide fixes from the BIRK / MRNA / 9618.HK / BA
workbook reviews, pinned with the case that motivated each.

D1 delete a zero-weight DCF; D2 dated mid-year discounting; D3 listed minorities at market;
D5 risk-on band; D6 cascaded margins fade; D7 a failed anchor degrades; D8 company-specific spread.
"""
import copy
import importlib.util
import io
import os
import tempfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

from src.agents.analysis import dcf_agent as d
from src.agents.analysis import guidance_forecast as gfm
from src.agents.analysis import three_statement as ts
from src.utils.valuation_workbook import build_workbook

_spec = importlib.util.spec_from_file_location("_wbtests", Path(__file__).resolve().parent / "test_valuation_workbook.py")
_wbt = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(_wbt)


# ── D2: dating ─────────────────────────────────────────────────────────────────

def test_birk_timing_counts_a_quarter_of_fy26_and_carries_to_the_valuation_date():
    tm = d._dcf_timing("2025-09-30", "2026-06-30", "2026-10-04")
    assert tm["flow_fractions"][0] == pytest.approx(92 / 365, abs=1e-3)       # Jul–Sep 2026 still to come
    assert tm["flow_fractions"][1:] == [1.0] * 9
    assert tm["roll_forward_years"] == pytest.approx(96 / 365.25, abs=1e-4)


def test_a_year_already_in_the_balance_sheet_contributes_no_flow():
    # FY0+1 ended before the balance-sheet date: its cash is in the bridge already.
    tm = d._dcf_timing("2024-12-31", "2026-03-31", "2026-04-15")
    assert tm["flow_fractions"][0] == 0.0
    assert 0.0 < tm["flow_fractions"][1] < 1.0


def test_undated_projection_is_the_legacy_year_end_projection():
    a = d._project_dcf(1000.0, 0.2, 0.05, 0.0, 0.09, 0.03, -0.05, 100.0, 10.0)
    b = d._project_dcf(1000.0, 0.2, 0.05, 0.0, 0.09, 0.03, -0.05, 100.0, 10.0, timing=None)
    assert a[0] == pytest.approx(b[0])


def test_dated_projection_is_mid_year_and_stubbed():
    tm = d._dcf_timing("2025-12-31", "2026-06-30", "2026-06-30")
    iv, pv_f, pv_t, rows = d._project_dcf(1000.0, 0.2, 0.0, 0.0, 0.10, 0.0, -0.05, 0.0, 1.0, timing=tm)
    f1 = rows[0]["flow_fraction"]
    assert rows[0]["fcf"] == pytest.approx(200.0 * f1)
    assert rows[0]["discount_factor"] == pytest.approx(1.1 ** (-f1 / 2))
    assert rows[1]["discount_factor"] == pytest.approx(1.1 ** (-(f1 + 0.5)))


def _dated_run():
    tm = d._dcf_timing("2025-09-30", "2026-06-30", "2026-10-04")
    sched = [0.18, 0.19, 0.2, 0.2, 0.21, 0.21, 0.21, 0.21, 0.21, 0.21]
    iv, pv_f, pv_t, rows = d._project_dcf(1000.0, 0.2, 0.08, 0.0, 0.09, 0.03, -0.05, 100.0, 10.0,
                                          margin_schedule=sched, minority_interest=40.0, preferred_equity=5.0,
                                          timing=tm)
    run = _wbt._run()
    dr = run["data"]["dcf_range"]["TEST"]
    for s in ("bear", "base", "bull"):
        legs = copy.deepcopy(dr[s]["leg_inputs"])
        legs["DCF"].update(value=iv, pv_fcf_per_share=pv_f, pv_tv_per_share=pv_t, projection_rows=rows,
                           growth_schedule=None, growth_base=0.08, timing=tm, minority_interest=40.0,
                           preferred_equity=5.0)
        dr[s]["leg_inputs"] = legs
        dr[s]["method_iv_table"]["DCF"] = iv
        dr[s]["intrinsic_value"] = round(0.4 * iv + 0.4 * 230.0 + 0.2 * 180.0, 2)
    iv_b = dr["base"]["intrinsic_value"]
    dr["12m_targets"] = {s: round(150.0 + 0.35 * (iv_b - 150.0), 2) for s in ("bear", "base", "bull")}
    return run, iv


def test_the_workbook_rebuilds_a_dated_projection_with_a_margin_path_to_zero():
    formulas = pytest.importorskip("formulas")
    run, iv = _dated_run()
    with tempfile.TemporaryDirectory() as td:
        p = os.path.join(td, "dated.xlsx")
        Path(p).write_bytes(build_workbook(run, "TEST"))
        sol = {k.upper(): v for k, v in formulas.ExcelModel().loads(p).finish().calculate().items()}
        ws = load_workbook(p)["DCF"]
        checks = [c.row for c in ws["A"] if c.value == "Check"]
        assert checks
        for r in checks:
            v = sol[f"'[DATED.XLSX]DCF'!B{r}".upper()]
            v = getattr(v, "value", v)[0][0]
            assert abs(v) < 1e-6, (r, v)


# ── 1A: the workbook ───────────────────────────────────────────────────────────

def test_d1_an_unweighted_dcf_has_no_tab_and_one_summary_line():
    run = _wbt._run_with({}, [{"method": "EV/EBITDA", "value_key": "EV/EBITDA", "bucket": "multi", "weight": 0.6},
                              {"method": "P/E", "value_key": "P/E", "bucket": "multi", "weight": 0.4}])
    wb = load_workbook(io.BytesIO(build_workbook(run, "TEST")))
    assert "DCF" not in wb.sheetnames
    lines = [str(c.value) for c in wb["Summary"]["A"] if c.value]
    assert sum(1 for x in lines if x.startswith("DCF EXCLUDED from the blend")) == 1


def test_capture_is_stated_with_its_reason():
    run = _wbt._run()
    run["data"]["dcf_range"]["TEST"]["pt_bridge"]["capture_reason"] = "35%: the standard rule"
    wb = load_workbook(io.BytesIO(build_workbook(run, "TEST")))
    assert any("35%: the standard rule" in str(c.value) for c in wb["Assumptions"]["A"] if c.value)


# ── 1C: the bridge ─────────────────────────────────────────────────────────────

def _row(**kw):
    base = {"net_debt": 100.0, "total_debt": 150.0, "cash_and_equivalents": 50.0, "short_term_investments": 0.0,
            "lease_liabilities": None, "period": "2025-12-31"}
    base.update(kw)
    return base


def test_us_gaap_leases_come_out_ifrs_leases_stay():
    us, b_us = d._valuation_net_debt(_row(lease_liabilities=30.0), "Consumer", "SBUX", "USD")
    ifrs, b_ifrs = d._valuation_net_debt(_row(lease_liabilities=30.0), "Consumer", "BIRK", "EUR")
    jd, _ = d._valuation_net_debt(_row(lease_liabilities=30.0), "Consumer", "09618.HK", "CNY")
    assert us == 70.0 and jd == 70.0 and ifrs == 100.0
    assert b_us["accounting_basis"] == "US GAAP" and b_ifrs["accounting_basis"] == "IFRS"


def test_biotech_treasury_is_cash_managed_care_float_is_not():
    r = _row(net_debt=-500.0, total_debt=1240.0, cash_and_equivalents=1740.0, short_term_investments=3200.0)
    biotech, _ = d._valuation_net_debt(dict(r), "Healthcare", "MRNA", "USD", "Biotechnology")
    plan, _ = d._valuation_net_debt(dict(r), "Healthcare", "XMCO", "USD", "Medical - Healthcare Plans")
    assert biotech == -3700.0
    # Plan IN1: a managed-care plan's cash and investments are regulated capital; with no parent-only
    # cash on record, net debt is the gross debt.
    assert plan == 1240.0


def test_interest_income_leaves_unlevered_fcf():
    from types import SimpleNamespace
    li = SimpleNamespace(report_period="2025-12-31", currency="CNY", revenue=1000.0, free_cash_flow=100.0,
                         interest_expense=10.0, interest_income=20.0, ebit=200.0, net_income=142.5)
    rows, _ = d._extract_annual_series([li])
    r = rows[-1]
    tax = r["ufcf_tax_rate"]
    assert r["ufcf_owner_earnings"] == pytest.approx(r["fcf_owner_earnings"] + 10.0 * (1 - tax) - 20.0 * (1 - tax))
    assert r["ufcf_interest_income_removed"] == pytest.approx(20.0 * (1 - tax))


def test_boeing_disposal_gain_is_not_core_earnings():
    prior = [{"operating_income": 1e9, "ebitda": 3.2e9, "depreciation_and_amortization": 2e9, "revenue": 80e9}] * 3
    ce = d._core_earnings({"operating_income": -5.42e9, "ebitda": 7.36e9, "depreciation_and_amortization": 1.95e9,
                           "revenue": 89.46e9, "net_income": 2.23e9, "ufcf_tax_rate": 0.15}, prior)
    assert ce["ebitda_core"] == pytest.approx(-3.47e9 + 0.2e9)        # the usual 0.2bn below the line stays
    assert ce["net_income_core"] < 0
    assert d._core_earnings({"operating_income": 10e9, "ebitda": 13.1e9, "depreciation_and_amortization": 3e9,
                             "revenue": 50e9, "net_income": 7e9}) == {}


def test_recurring_associate_income_is_core():
    # Sembcorp-shaped: a large share of associates' profit below the operating line, every year.
    rows = [{"operating_income": 0.5e9, "ebitda": 1.8e9, "depreciation_and_amortization": 0.3e9, "revenue": 7e9}] * 4
    assert d._core_earnings({"operating_income": 0.55e9, "ebitda": 1.9e9, "depreciation_and_amortization": 0.3e9,
                             "revenue": 7.2e9, "net_income": 1.0e9}, rows) == {}


# ── 1B / 1E: the guidance forecast ─────────────────────────────────────────────

_HIST = {"revenue": 1000.0, "ebit": 200.0, "ebit_margin": 0.20, "da": 60.0, "da_pct_revenue": 0.06, "capex": 70.0,
         "net_income": 140.0, "interest": 10.0, "shares": 100.0, "tax_rate": 0.25, "tax_rate_source": "history",
         "capex_alpha": 0.1, "capex_alpha_n": 3, "nwc_intensity": 0.05, "nwc_n": 3, "buyback_median": 0.0,
         "roic_median": 0.15, "years": 5, "fy0": 2025}


def _fc(block, **kw):
    return gfm.build_forecast(block, scenario="base", series=None, hist=dict(_HIST), profile_name="Generic", sector="Consumer",
                              wacc=0.09, tgr=0.025, shares=100.0, net_debt=0.0, **kw)


def test_an_ebitda_margin_endpoint_is_converted_to_ebit():
    fc = _fc({"fiscal_year_1": "FY2026", "estimates": {"base": {"revenue_growth_fy1": 0.10, "ebitda_margin_fy1": 0.30}}})
    assert fc["margin_target"] == pytest.approx(0.24)


def test_research_fy1_a_year_late_runs_the_engine_path_first():
    fc = _fc({"fiscal_year_1": "FY2027", "estimates": {"base": {"revenue_growth_fy1": 0.03, "revenue_growth_fy2": 0.04}}},
             engine_growth_path=[-0.02] * 10)
    assert fc["growth_schedule"][:3] == pytest.approx([-0.02, 0.03, 0.04])
    assert fc["rows"][0]["phase"] == "engine path"


def test_guided_eps_is_converted_into_the_valuation_currency():
    blk = {"fiscal_year_1": "FY2026", "guidance": {"eps": {"currency": "EUR"}},
           "estimates": {"base": {"revenue_growth_fy1": 0.10, "eps_fy1": 1.0}}}
    a = _fc(blk)
    b = _fc(blk, fx_to_valuation=1.2, valuation_currency="USD")
    assert b["deconstruction"]["eps_T_guided"] == pytest.approx(1.2 * a["deconstruction"]["eps_T_guided"])


def test_three_statements_are_built_in_the_statement_currency():
    fc = {"rows": [{"revenue": 1120.0, "growth": 0.0, "ebit": 112.0, "da": 56.0, "capex": 56.0, "delta_nwc": 0.0}],
          "history": {"tax_rate": 0.25}}
    op = {"fiscal_year": 2025, "revenue": 1000.0, "cash": 100.0, "equity": 400.0, "ppe": 300.0, "payables": 0.0}
    out = ts.build(fc, op, {}, years=1, fx=1.12)
    assert out["income"]["revenue"][0] == pytest.approx(1000.0)
    assert out["income"]["growth"][0] == pytest.approx(0.0)


# ── 1F: blend and peers ────────────────────────────────────────────────────────

def test_d7_a_failed_anchor_moves_its_weight_to_the_dcf_and_says_degraded():
    methods = [{"name": "EV/EBIT (norm)", "weight": 0.3, "anchor": True},
               {"name": "DCF", "weight": 0.2}, {"name": "P/E", "weight": 0.5}]
    flags: list = []
    iv, bd = d._blend_methods(methods, {"EV/EBIT (norm)": None, "DCF": 50.0, "P/E": 100.0}, 1.0, flags, 0.5)
    assert bd["anchor_degraded"]["moved_to"] == "the DCF"
    assert iv == pytest.approx((0.5 * 50.0 + 0.5 * 100.0))
    assert any(f.startswith("DEGRADED") for f in flags)


def test_the_valued_company_is_not_in_its_own_curated_basket(monkeypatch):
    from src.data import regional_comps as rc
    seen = {}

    def _q(sql, params):
        seen["params"] = params
        return []
    monkeypatch.setattr(rc._db, "query", _q)
    monkeypatch.setattr(rc, "_ensure_table", lambda: None)
    rc.basket_multiples("HKSE", ("09988.HK", "00700.HK"), "China Internet Platform", exclude="09988.HK")
    assert "9988.HK" not in seen["params"] and "0700.HK" in seen["params"]


# ── Phase 1b (MOH / JD / D05 / 9988 reviews) ───────────────────────────────────

def test_en2_recovery_discount_prices_the_excess_when_consensus_reaches_it():
    v, info = d._recovery_discounted(1220.0, [274.0, 532.0, 700.0, 900.0, 1240.0], 0.08)
    assert info["years_to_recover"] == 5
    assert v == pytest.approx(274.0 + (1220.0 - 274.0) / 1.08 ** 5)
    assert d._recovery_discounted(200.0, [274.0], 0.08) == (200.0, None)          # below forward: unchanged


def test_en3_an_incompatible_margin_gives_way_to_guided_eps():
    blk = {"fiscal_year_1": "FY2026", "estimates": {"base": {"revenue_growth_fy1": 0.05, "ebitda_margin_fy1": 0.30, "eps_fy1": 0.5}}}
    fc = _fc(blk)
    dec = fc["deconstruction"]
    assert "margin_T_conflicting" in dec and dec["margin_T_guided"] < 0.10
    assert any("guided EPS implies" in f or "not compatible" in f for f in fc["flags"])


def test_in2_bank_scenarios_move_drivers_not_the_answer():
    import inspect
    src = inspect.getsource(d._compute_method_value)
    assert "_BANK_SCENARIO_ROE_SHIFT" in src and "_BANK_SCENARIO_COE_SHIFT" in src


def test_managed_care_regulated_cash_matches_the_engines_sector_name():
    # Molina resolves to "HealthcareServices"; the guard keyed on "Healthcare" never fired.
    r = {"net_debt": -4000.0, "total_debt": 4000.0, "cash_and_equivalents": 5000.0, "short_term_investments": 3000.0,
         "lease_liabilities": 184.0, "period": "2025-12-31"}
    nd, b = d._valuation_net_debt(dict(r), "HealthcareServices", "XMCO", "USD", "Medical - Healthcare Plans")
    assert b["regulated_cash_excluded"] is True and nd == 4000.0 - 184.0


def test_an_accepted_parent_cash_entry_counts_and_a_proposed_one_does_not(monkeypatch):
    from src.data import valuation_constants as vc
    doc = {"parent_cash": {"entries": {"MOH": {"value": 290e6, "status": "ACCEPTED"}, "XMCO": {"value": 1e9, "status": "PROPOSED"}}}}
    monkeypatch.setattr(vc, "load", lambda *a, **k: doc)
    assert d._parent_cash("MOH") == 290e6 and d._parent_cash("XMCO") is None


# ── Visa / Vertex reviews (owner, 2026-10-04): EV2, EV7, IV1, IV3 ───────────────

def test_ev2_routine_insider_selling_does_not_move_wacc():
    # Visa: $83m of sales (~0.01% of a ~$700bn cap) with a CEO/CFO conviction-sell flag added +16bp.
    bps, _ = d._insider_wacc_modifier({"net_buying_12m_usd": -83e6, "gross_sell_value_12m": 83e6,
                                       "conviction_sell_flag": True}, 700e9)
    assert bps == 0.0
    # Material selling (1% of cap) still widens.
    bps2, _ = d._insider_wacc_modifier({"net_buying_12m_usd": -7e9, "gross_sell_value_12m": 7e9}, 700e9)
    assert bps2 > 0
    # Buying is unaffected by the selling threshold.
    bps3, _ = d._insider_wacc_modifier({"net_buying_12m_usd": 1.4e9, "gross_buy_value_12m": 1.4e9}, 700e9)
    assert bps3 < 0


def test_ev7_terminal_multiple_band_is_two_sided_and_roic_aware():
    import inspect
    src = inspect.getsource(gfm)
    assert "ROIC-justified" in src and "BELOW the floor" in src


def test_iv1_payment_networks_take_the_unlevered_basis():
    assert d._interest_is_cost_of_goods("Payment Networks", "Financials") is False
    assert d._interest_is_cost_of_goods("Asset Manager", "Financials") is True
    assert d._interest_is_cost_of_goods("Money Center Bank", "Financials") is True


def test_iv3_biotech_long_term_securities_are_cash_pharma_stakes_are_not():
    r = _row(net_debt=-500.0, total_debt=1240.0, cash_and_equivalents=1740.0, short_term_investments=3200.0,
             long_term_investments=10000.0)
    biotech, b = d._valuation_net_debt(dict(r), "Healthcare", "VRTX", "USD", "Biotechnology")
    pharma, _ = d._valuation_net_debt(dict(r), "Healthcare", "PFE", "USD", "Drug Manufacturers - General")
    assert biotech == -13700.0 and b["long_term_investments_netted"] == 10000.0
    assert pharma == -3700.0


def test_iv2_pipeline_is_an_add_on_for_revenue_stage_drug_profiles():
    assert "Commercial Biotech" in d._PIPELINE_ADDON_PROFILES
    assert "Pre-approval Biotech" not in d._PIPELINE_ADDON_PROFILES
    # An approved product's sales are in the operating legs' revenue: the add-on carries unapproved assets only.
    import inspect
    assert 'a.get("phase") != "approved"' in inspect.getsource(d)


def test_ev6_ntm_roll_on_the_guidance_overlay():
    # Vertex, October 2026: FY2026 EPS $17.45 is three-quarters gone; NTM weights FY2027's $20.50 at 75%.
    est = {"confidence": "HIGH",
           "estimates": {"base": {"eps_fy1": 17.45, "eps_fy2": 20.50, "revenue_growth_fy1": 0.10,
                                  "revenue_growth_fy2": 0.10, "ebitda_margin_fy1": 0.45, "ebitda_margin_fy2": 0.45}}}
    fwd = {"eps": {"base": 18.0}, "ebitda": {"base": 1.0}, "revenue": {"base": 1.0}, "ebit": {"base": 1.0}}
    o0 = d._guidance_forward_overlay(est, copy.deepcopy(fwd), "base", None, 1000.0)
    o75 = d._guidance_forward_overlay(est, copy.deepcopy(fwd), "base", None, 1000.0, ntm_e=0.75)
    assert o0["eps"]["base"] == pytest.approx(17.45)
    assert o75["eps"]["base"] == pytest.approx(0.25 * 17.45 + 0.75 * 20.50)
    assert o75["revenue"]["base"] == pytest.approx(1100.0 * (0.25 + 0.75 * 1.10))


def test_bridge_adjustments_price_only_when_accepted(monkeypatch):
    from src.data import valuation_constants as vc
    doc = {"bridge_adjustments": {"entries": {"V": {"debt_like": 363.0, "status": "ACCEPTED"},
                                              "MA": {"debt_like": 999.0, "status": "PROPOSED"}}}}
    monkeypatch.setattr(vc, "load", lambda *a, **k: doc)
    v, b = d._valuation_net_debt(_row(), "Financials", "V", "USD")
    ma, _ = d._valuation_net_debt(_row(), "Financials", "MA", "USD")
    assert v == 463.0 and b["debt_like_items"] == 363.0 and ma == 100.0


def test_the_visa_entry_is_owner_accepted():
    from src.data import valuation_constants as vc
    e = vc.load()["bridge_adjustments"]["entries"]["V"]
    assert e["status"] == "ACCEPTED" and e["preferred_in_shares"] is True
    assert d._bridge_adjustment("V")["shares_as_converted"] == 1880000000.0


def test_d8_scenario_wacc_shift_is_removed():
    assert set(d._WACC_SCENARIO_SHIFT.values()) == {0.0}


def test_payment_networks_price_on_the_toll_road_basket():
    from src.data.regional_comps import PROFILE_PEER_BASKETS
    assert PROFILE_PEER_BASKETS["Payment Networks"]["US"] == ("V", "MA", "AXP", "SPGI", "MCO", "ICE", "CME", "MSCI")


def test_guidance_eps_on_the_forward_legs_is_converted_into_the_valuation_currency():
    # IHH (2026-10-06): RM0.26 / RM0.29 guided EPS sat unconverted in an SGD forward P/E (S$5.09 vs ~S$1.59).
    est = {"confidence": "HIGH", "estimates": {"base": {"eps_fy1": 0.26, "eps_fy2": 0.29}}}
    fwd = {"eps": {"base": 0.1}, "ebitda": {"base": 1.0}, "revenue": {"base": 1.0}, "ebit": {"base": 1.0}}
    o = d._guidance_forward_overlay(est, copy.deepcopy(fwd), "base", None, None, ntm_e=0.76,
                                    fx_to_valuation=0.31288, valuation_currency="SGD")
    assert o["eps"]["base"] == pytest.approx((0.24 * 0.26 + 0.76 * 0.29) * 0.31288)
    est["guidance"] = {"eps": {"currency": "SGD"}}
    o2 = d._guidance_forward_overlay(est, copy.deepcopy(fwd), "base", None, None, fx_to_valuation=0.31288,
                                     valuation_currency="SGD")
    assert o2["eps"]["base"] == pytest.approx(0.26)


def test_the_subject_exclusion_and_frozen_members_reach_the_payload(monkeypatch):
    # 2026-10-06: computed in get_regional_multiples, dropped in get_sector_peer_multiples' basis copy.
    from src.data import sector_profiles as sp
    row = {"value": 46.36, "basis": "industry", "cohort": "large", "peer_count": 5, "key": "Biotechnology",
           "exchange": "HKSE", "subject_excluded": True, "value_with_subject": 52.91,
           "members_used": [{"symbol": "6160.HK", "value": 40.0}]}
    import inspect
    src = inspect.getsource(sp.get_sector_peer_multiples)
    assert '"subject_excluded", "members_used", "excluded_no_market_cap", "value_with_subject"' in src
    tr = d._multiples_trace({"pe": 46.36, "_comp_basis": {"pe": row}})
    assert tr["fields"]["pe"]["subject_excluded"] is True and tr["fields"]["pe"]["members_used"]


def test_e26_holders_and_institutions_are_not_insiders():
    from src.agents.intelligence.insider_activity_agent import _is_reportable_insider as rep
    assert rep("10% Owner", "LILLY ENDOWMENT INC") is False
    assert rep(None, "Lilly Endowment Inc") is False
    assert rep("Director, 10% Owner", "Jane Doe") is True
    assert rep("EVP & CFO", "John Smith") is True


def test_e16_cash_runway_is_computed_and_overrides_the_extractor(monkeypatch):
    from src.data import sector_kpi_framework as k
    monkeypatch.setattr(k, "_fmp_risk_kpis", lambda t: {"cash_runway_qtrs": 16.6, "cash_runway_years": 4.15})
    out = k._augment_metrics_with_fmp_risk("MRNA", {"cash_runway_qtrs": 1.7})
    assert out["cash_runway_qtrs"] == 16.6


def test_e24_a_research_fy1_without_fy2_rolls_on_consensus_growth_and_says_so():
    est = {"confidence": "HIGH", "estimates": {"base": {"eps_fy1": 36.0}}}
    fwd = {"eps": {"base": 40.0}, "ebitda": {"base": 1.0}, "revenue": {"base": 1.0}, "ebit": {"base": 1.0},
           "_fy1_fy2": {"eps": {"base": (34.0, 44.2)}}}
    o = d._guidance_forward_overlay(est, copy.deepcopy(fwd), "base", None, None, ntm_e=0.76)
    assert o["eps"]["base"] == pytest.approx(0.24 * 36.0 + 0.76 * 36.0 * 44.2 / 34.0)
    assert "consensus FY+2/FY+1 growth" in o["_source"]["eps"]["base"]
    o2 = d._guidance_forward_overlay(est, {**copy.deepcopy(fwd), "_fy1_fy2": {}}, "base", None, None, ntm_e=0.76)
    assert o2["eps"]["base"] == pytest.approx(36.0) and "no FY+2 to roll" in o2["_source"]["eps"]["base"]


def test_e23_the_forward_growth_adjustment_is_the_gap_over_peers_bounded():
    fc = {"_fy1_fy2": {"revenue": {"base": (100.0, 125.0)}}}
    f, note = d._forward_growth_adjustment(fc, {"growth_avg": 0.05})
    assert f == pytest.approx(1.3) and "held at" in note             # (1.25/1.05)^3 = 1.69 -> the owner band 1.30
    f2, _ = d._forward_growth_adjustment({"_fy1_fy2": {"revenue": {"base": (100.0, 105.0)}}}, {"growth_avg": 0.05})
    assert f2 == pytest.approx(1.0)
    assert d._forward_growth_adjustment({}, {"growth_avg": 0.05})[0] is None


def test_e28_adjusted_eps_neither_sets_the_margin_nor_fails_the_share_check():
    import inspect
    src = inspect.getsource(gfm)
    assert 'block.get("_eps_adjusted")' in src and "n/a: guided EPS is on an adjusted (non-GAAP) basis" in src


def test_e1_windfall_years_leave_normalised_earnings_for_non_cyclicals():
    # Pfizer: 2021-22 COVID years at 27% / 31% net margin against ~12% otherwise.
    rows = [{"period": f"{y}-12-31", "revenue": r, "net_income": n} for y, r, n in
            ((2021, 81.3e9, 22.0e9), (2022, 100.3e9, 31.4e9), (2023, 58.5e9, 2.1e9), (2024, 63.6e9, 8.0e9), (2025, 62.6e9, 7.8e9))]
    assert d._windfall_periods(rows, "net_income") == ["2021", "2022"]
    full = d._normalized_earnings(rows, "net_income")
    ex = d._normalized_earnings(rows, "net_income", exclude_windfalls=True)
    assert ex < full and ex / 62.6e9 == pytest.approx((2.1 / 58.5 + 8.0 / 63.6 + 7.8 / 62.6) / 3, rel=1e-6)


def test_ddm_low_yield_and_scale_cap_rules_are_wired():
    import inspect
    src = inspect.getsource(d)
    assert d._DDM_MIN_YIELD == 0.02 and "if _yld < _DDM_MIN_YIELD or _bb_share >= _DDM_MAX_BUYBACK_SHARE:" in src
    assert "_cons_near_ok = True" in src and "the capped rate governs year 3 onward" in src


def test_e30_preferreds_and_notes_are_not_comps():
    from src.data.regional_comps import is_non_common_security as f
    assert f("Brookfield Finance Inc. 4.625% Subordinated Notes due 2080", "BNH")
    assert f("Strive, Inc. Variable Rate Series A Perpetual Preferred Stock", "SATA")
    assert f("Wells Fargo & Company", "WFC-PL")
    assert not f("Brookfield Corporation", "BN") and not f("Noteworthy AI", "NOTE")


def test_e27_the_risk_on_floor_only_undoes_the_cut_and_the_wacc_tab_shows_it():
    import inspect
    src = inspect.getsource(d)
    assert "_floor = round(min(_capm_wacc - _RISK_ON_CAPM_BAND, wacc - _ov), 4)" in src
    from src.utils import valuation_workbook as vw
    assert "CAPM band: risk-on floor (D5)" in inspect.getsource(vw)


def test_i9_i10_i11_curated_baskets(monkeypatch):
    from src.data import regional_comps as rc
    assert "REGN" in rc.PROFILE_PEER_BASKETS["Commercial Biotech"]["US"]
    assert rc.TICKER_BASKET_BLENDS["BLK"] == (("Asset Manager", 0.75), ("Alt Asset Manager", 0.25))
    assert "SES" in rc.CROSS_MARKET_BASKETS["Grocery & Discount Retail"]["markets"]
    calls = []
    monkeypatch.setattr(rc, "basket_multiples", lambda ex, syms, key, **k: calls.append((ex, key)) or
                        {"pe_ntm": {"value": 10.0 if key == "Asset Manager" else 20.0, "peer_count": 6, "members": []}})
    out = rc.profile_basket_multiples("US", "Asset Manager", exclude="BLK")
    assert out["pe_ntm"]["value"] == pytest.approx(0.75 * 10 + 0.25 * 20)
    rc.profile_basket_multiples("SES", "Grocery & Discount Retail", exclude="OV8.SI")
    assert calls[-1] == ("XMKT", "Grocery & Discount Retail")


def test_market_minority_reads_consolidated_stakes_below_half(monkeypatch):
    # Owner, 2026-10-07 (IHH review): Fortis at 31% is consolidated -- a `consolidated: true` entry prices its
    # outside holders at market; a sub-50% entry without the flag is an associate and prices nothing.
    from src.agents.analysis import holdco_sotp as hs
    from src.data import valuation_constants as vc
    reg = {"listed_subsidiaries": {"entries": {
        "IHHX": [{"listed": "FORTIS.NS", "stake_pct": 0.3117, "consolidated": True, "status": "ACCEPTED"},
                 {"listed": "ASSOC.NS", "stake_pct": 0.30, "status": "ACCEPTED"}],
        "PROP": [{"listed": "C2PU.SI", "stake_pct": 0.3294, "consolidated": True, "status": "PROPOSED"}]}}}
    monkeypatch.setattr(vc, "load", lambda: reg)
    monkeypatch.setattr(hs, "_market_value", lambda listed, end: 1000.0)
    monkeypatch.setattr(hs, "_fx", lambda a, b: 0.05)
    monkeypatch.setattr(hs, "template_for", lambda t: None)
    out = hs.listed_minority_at_market("IHHX", "2026-10-07", "MYR")
    assert [p["listed"] for p in out["parts"]] == ["FORTIS.NS"]
    assert abs(out["value"] - (1 - 0.3117) * 1000.0 * 0.05) < 1e-9
    assert hs.listed_minority_at_market("PROP", "2026-10-07", "MYR") is None   # PROPOSED prices nothing


def test_blk_bridge_entry_and_minority_override_are_wired():
    # Owner, 2026-10-07 (BlackRock review): the bridge entry prices only once ACCEPTED, and its
    # minority_interest replaces the feed's (CIP NCI and Subco units are not outside claims).
    import inspect
    from src.agents.analysis import dcf_agent as d
    from src.data import valuation_constants as vc
    e = vc.load()["bridge_adjustments"]["entries"]["BLK"]
    assert e["minority_interest"] < 1e9 and e["shares_as_converted"] == 164_600_000.0
    if str(e.get("status")).upper() != "ACCEPTED":
        assert d._bridge_adjustment("BLK") == {}
    src = inspect.getsource(d)
    assert '_bridge_adjustment(ticker).get("minority_interest")' in src


def test_rerun_a_failed_share_count_check_prices_the_forecast_eps():
    # Sample rerun 2026-10-07 (REGN): research EPS $35.5 against forecast NI that implies $41.5 on today's
    # shares -- invariant 3 fails, and the forward P/E prices the forecast's EPS, labelled as such.
    est = {"confidence": "HIGH", "estimates": {"base": {"eps_fy1": 32.0, "eps_fy2": 35.5}}}
    fwd = {"eps": {"base": 60.0}, "ebitda": {"base": 1.0}, "revenue": {"base": 1.0}, "ebit": {"base": 1.0}}
    gf = {"rows": [{"eps": 39.0, "ebit": 5.0}, {"eps": 41.5}], "invariants": [{"id": 3, "ok": False}]}
    o = d._guidance_forward_overlay(est, copy.deepcopy(fwd), "base", gf, None, ntm_e=0.76)
    assert o["eps"]["base"] == pytest.approx(0.24 * 39.0 + 0.76 * 41.5)
    assert "failed the share-count check" in o["_source"]["eps"]["base"]
    gf_ok = {**gf, "invariants": [{"id": 3, "ok": True}]}
    o2 = d._guidance_forward_overlay(est, copy.deepcopy(fwd), "base", gf_ok, None, ntm_e=0.76)
    assert o2["eps"]["base"] == pytest.approx(0.24 * 32.0 + 0.76 * 35.5)


def test_rerun_e29_collaboration_heavy_revenue_stands_the_ev_revenue_leg_down():
    # Sample rerun 2026-10-07 (REGN): a blended leg values the whole company, so the revenue leg stands down
    # above 30% collaboration income instead of pricing the product share alone; the E25 flag leaves when
    # the guidance forecast replaces the schedule.
    import inspect
    src = inspect.getsource(d)
    assert "fwd_rev = fwd_rev * (1.0 - _collab_share)" not in src
    assert '"_ev_rev_collab_stood_down"' in src and src.count("_ev_rev_collab_stood_down") >= 3
    assert 'startswith("Consensus sets DCF years 1-2")' in src


# ── REGN review (2026-10-07): EPS basis, LOE, the effective method ────────────

def test_regn_forward_pe_reverts_to_street_eps_in_every_scenario_when_guidance_diverges():
    def sc(s, cur, cons):
        return {"eps": {s: cur}, "_consensus": {"eps": {s: cons}}, "_source": {"eps": {s: "forecast EPS"}}}
    out, dec, flag = d._eps_street_basis_guard(sc("base", 42.1, 60.4), "base", None, 15)
    assert dec is True and out["eps"]["base"] == 60.4 and "street (adjusted) EPS" in flag
    out_b, dec_b, flag_b = d._eps_street_basis_guard(sc("bear", 29.9, 45.7), "bear", dec, 15)
    assert out_b["eps"]["bear"] == 45.7 and flag_b is None                       # one basis for all scenarios
    _, dec2, _ = d._eps_street_basis_guard(sc("base", 43.7, 44.5), "base", None, 15)   # LLY: in line, kept
    assert dec2 is False
    _, dec3, _ = d._eps_street_basis_guard(sc("base", 42.1, 60.4), "base", None, 3)    # thin coverage, kept
    assert dec3 is False


def test_i14_loe_overlay_erodes_named_drugs_and_haircuts_the_terminal_value():
    entry = {"total_revenue": 100.0, "drugs": [
        {"name": "Dupi", "revenue_fy": 40.0, "loe_year": 2031, "modality": "biologic"},
        {"name": "Late", "revenue_fy": 20.0, "loe_year": 2040, "modality": "small_molecule"}]}
    o = d._franchise_loe_overlay(entry, 2025, [0.05] * 10, 0.08, 0.025)
    idx = o["index"]
    assert idx[4] == pytest.approx(1.0)                                         # 2030: before any LOE
    # 2031 (LOE year): Dupi frozen at its 2030 level (no company growth after LOE) and down 15% on the curve
    # owner 2026-10-07: the multi-franchise 50% replacement offsets the in-horizon erosion too
    assert idx[5] == pytest.approx(1.0 - 0.5 * (0.4 - 0.4 * 0.85 / 1.05), abs=1e-5)
    assert idx[9] == pytest.approx(1.0 - 0.5 * (0.4 - 0.4 * 0.35 / 1.05 ** 5), abs=1e-5)  # 2035: 65% lost, frozen since 2030
    assert (1 + o["growth_schedule"][5]) == pytest.approx(1.05 * idx[5] / idx[4])
    # after the horizon: Dupi is fully eroded by 2035; Late loses 90% from 2040, five years past 2035, discounted
    disc = 1.025 / 1.08
    lost = 0.2 * 0.90 * disc ** 5 * (1 - 0.5)          # multi-franchise (40% + 20%): half credited to the pipeline
    assert o["terminal_replacement"] == 0.5
    assert o["terminal_multiplier"] == pytest.approx(1.0 - lost / idx[9], abs=1e-5)
    # a drug already eroding loses only what is left of its curve
    e2 = {"total_revenue": 100.0, "drugs": [{"name": "Eylea", "revenue_fy": 30.0, "loe_year": 2024, "modality": "biologic"}]}
    o2 = d._franchise_loe_overlay(e2, 2025, [0.0] * 10, 0.08, 0.025)
    # years 1-2 (2026-27) are guidance / consensus and already price it; 2028 = year 5 vs year 4 at end-2027
    assert o2["index"][:2] == [1.0, 1.0]
    assert o2["index"][2] == pytest.approx(1.0 - 0.7 * 0.3 * (0.65 - 0.60) / (1 - 0.60), abs=1e-5)   # single franchise: 30% replaced
    assert d._franchise_loe_entry("REGN") == {} or d._franchise_loe_entry("REGN").get("status") == "ACCEPTED"


def test_i14_terminal_multiplier_scales_the_terminal_value_and_the_workbook_reproduces_it():
    a = d._project_dcf(1000.0, 0.2, 0.05, 0.0, 0.09, 0.03, -0.05, 100.0, 10.0)
    b = d._project_dcf(1000.0, 0.2, 0.05, 0.0, 0.09, 0.03, -0.05, 100.0, 10.0, terminal_multiplier=0.8)
    assert b[2] == pytest.approx(0.8 * a[2]) and b[1] == pytest.approx(a[1])
    formulas = pytest.importorskip("formulas")
    tm = d._dcf_timing("2025-09-30", "2026-06-30", "2026-10-04")
    sched = [0.18, 0.19, 0.2, 0.2, 0.21, 0.21, 0.21, 0.21, 0.21, 0.21]
    iv, pv_f, pv_t, rows = d._project_dcf(1000.0, 0.2, 0.08, 0.0, 0.09, 0.03, -0.05, 100.0, 10.0,
                                          margin_schedule=sched, minority_interest=40.0, preferred_equity=5.0,
                                          timing=tm, terminal_multiplier=0.8)
    run = _wbt._run()
    dr = run["data"]["dcf_range"]["TEST"]
    for s in ("bear", "base", "bull"):
        legs = copy.deepcopy(dr[s]["leg_inputs"])
        legs["DCF"].update(value=iv, pv_fcf_per_share=pv_f, pv_tv_per_share=pv_t, projection_rows=rows,
                           growth_schedule=None, growth_base=0.08, timing=tm, minority_interest=40.0,
                           preferred_equity=5.0, terminal_loe_multiplier=0.8)
        dr[s]["leg_inputs"] = legs
        dr[s]["method_iv_table"]["DCF"] = iv
        dr[s]["intrinsic_value"] = round(0.4 * iv + 0.4 * 230.0 + 0.2 * 180.0, 2)
    iv_b = dr["base"]["intrinsic_value"]
    dr["12m_targets"] = {s: round(150.0 + 0.35 * (iv_b - 150.0), 2) for s in ("bear", "base", "bull")}
    with tempfile.TemporaryDirectory() as td:
        p = os.path.join(td, "loe.xlsx")
        Path(p).write_bytes(build_workbook(run, "TEST"))
        sol = {k.upper(): v for k, v in formulas.ExcelModel().loads(p).finish().calculate().items()}
        ws = load_workbook(p)["DCF"]
        checks = [c.row for c in ws["A"] if c.value == "Check"]
        assert checks
        for r in checks:
            v = sol[f"'[LOE.XLSX]DCF'!B{r}".upper()]
            v = getattr(v, "value", v)[0][0]
            assert abs(v) < 1e-6, (r, v)


def test_regn_a_quarantined_pipeline_says_the_valuation_is_the_commercial_business_alone():
    import inspect
    assert "not the hybrid SOTP the method selection chose" in inspect.getsource(d)


def test_i14_big_pharma_terminal_credits_half_the_post_horizon_loss_to_the_pipeline():
    # Owner, 2026-10-07: 50% replacement for big pharma; every other profile takes the whole loss.
    entry = {"total_revenue": 100.0, "drugs": [{"name": "Tirz", "revenue_fy": 56.0, "loe_year": 2036, "modality": "small_molecule"}]}
    full = d._franchise_loe_overlay(entry, 2025, [0.05] * 10, 0.08, 0.025, profile_name="Commercial Biotech")
    half = d._franchise_loe_overlay(entry, 2025, [0.05] * 10, 0.08, 0.025, profile_name="Big Pharma (Consolidated DCF)")
    assert half["terminal_replacement"] == 0.5 and full["terminal_replacement"] == 0.30   # single franchise: 30% tier
    assert (1 - half["terminal_multiplier"]) == pytest.approx((0.5 / 0.7) * (1 - full["terminal_multiplier"]), abs=1e-5)
    from src.data import valuation_constants as vc
    assert vc.load()["franchise_loe"]["entries"]["REGN"]["status"] == "ACCEPTED"
    assert d._franchise_loe_entry("REGN")["total_revenue"] == 14342900000.0
    assert d._franchise_loe_entry("LLY")["status"] == "ACCEPTED"


def test_i14_tiered_replacement_and_the_pipeline_double_count_guard():
    # Owner, 2026-10-07: big pharma 50%; other drug developers 30% single-mechanism, 50% with multi-franchise
    # proof (largest franchise <= 50%, two at 10%+, drugs grouped by INN); an accepted pipeline's risk-adjusted
    # peak replaces part of the loss explicitly, and the terminal credit covers only the rest.
    from src.data import valuation_constants as vc
    reg = vc.load()["franchise_loe"]["entries"]
    assert d._loe_replacement_tier(reg["VRTX"], "Commercial Biotech")[0] == 0.30
    assert d._loe_replacement_tier(reg["ALNY"], "Commercial Biotech")[0] == 0.30
    assert d._loe_replacement_tier(reg["REGN"], "Commercial Biotech")[0] == 0.50       # aflibercept rows grouped
    assert d._loe_replacement_tier(reg["LLY"], "Big Pharma (Consolidated DCF)")[0] == 0.50
    entry = {"total_revenue": 100.0, "drugs": [{"name": "CF", "inn": "x", "revenue_fy": 90.0, "loe_year": 2038, "modality": "small_molecule"}]}
    base = d._franchise_loe_overlay(entry, 2025, [0.0] * 10, 0.08, 0.025, profile_name="Commercial Biotech")
    lost = 0.9 * 0.9 * 100.0                                                         # nominal, flat growth
    g = d._franchise_loe_overlay(entry, 2025, [0.0] * 10, 0.08, 0.025, profile_name="Commercial Biotech",
                                 pipeline_ra_peak=0.1 * lost)
    assert base["terminal_replacement"] == pytest.approx(0.30) and g["terminal_replacement"] == pytest.approx(0.20)
    assert "already priced by the accepted pipeline" in g["terminal_replacement_basis"]
    assert d._pipeline_risk_adjusted_peak_usd([{"phase": "phase_3", "peak_sales_usd": 100.0, "ptrs_override": 0.5},
                                               {"phase": "approved", "peak_sales_usd": 999.0},
                                               {"phase": "filed", "peak_sales_usd": 10.0}]) == pytest.approx(58.5)
    for t in ("600276.SS", "01276.HK", "01801.HK", "09688.HK"):
        assert reg[t]["caveat"] and reg[t]["status"] == "ACCEPTED"


# ── REGN review 3 (2026-10-07): named peers, street consensus, net-debt bridge, one revenue path ───────

def test_regn3_a_curated_basket_records_its_named_members_and_values(monkeypatch):
    from src.data import regional_comps as rc
    rows = [{"symbol": s_, "name": n_, "market_cap": mc, "computed_at": "2026-10-07T00:00:00",
             "metrics_json": __import__("json").dumps({"pe_ntm": {"value": v, "in_band": True}})}
            for s_, n_, mc, v in (("AMGN", "Amgen", 1.8e11, 14.0), ("GILD", "Gilead", 1.4e11, 15.0),
                                  ("VRTX", "Vertex", 1.2e11, 22.0), ("BIIB", "Biogen", 2.5e10, 9.0),
                                  ("ALNY", "Alnylam", 3.0e10, 40.0), ("REGN", "Regeneron", 8.0e10, 16.0))]
    monkeypatch.setattr(rc, "_ensure_table", lambda: None)
    monkeypatch.setattr(rc._db, "query", lambda sql, args: [r_ for r_ in rows if r_["symbol"] in args[1:]])
    monkeypatch.setattr(rc, "_age_days", lambda s_: 0.0)
    out = rc.basket_multiples("US", ("AMGN", "GILD", "VRTX", "BIIB", "ALNY", "REGN"), "Commercial Biotech", exclude="REGN")
    cell = out["pe_ntm"]
    assert cell["value"] == pytest.approx(15.0) and cell.get("subject_excluded") is True
    named = {m["symbol"]: (m["name"], m["value"]) for m in cell["members_used"]}
    assert named == {"AMGN": ("Amgen", 14.0), "GILD": ("Gilead", 15.0), "VRTX": ("Vertex", 22.0),
                     "BIIB": ("Biogen", 9.0), "ALNY": ("Alnylam", 40.0)}


def test_regn3_workbook_shows_ntm_peers_street_consensus_and_the_net_debt_bridge():
    formulas = pytest.importorskip("formulas")
    run = _wbt._run()
    dr = run["data"]["dcf_range"]["TEST"]
    mem = [{"symbol": s_, "name": s_, "market_cap": 1e10, "value": v} for s_, v in (("A", 14.0), ("B", 15.0), ("C", 22.0))]
    dr["multiples_used"] = {"comp_market": "US", "fields": {"pe_ntm": {"value": 15.0, "basis": "profile", "cohort": "all",
                                                                       "key": "Commercial Biotech", "exchange": "US", "members_used": mem}}}
    dr["street_consensus"] = {"source": "FMP analyst estimates", "fy1_period": "2026-12-31", "fy2_period": "2027-12-31",
                              "eps": {"fy1": 55.94, "fy2": 61.85, "ntm": 60.47}, "revenue": {"fy1": 1.7e10, "fy2": 1.86e10, "ntm": 1.83e10},
                              "analyst_count_eps": 15}
    nd_used = (dr["base"]["leg_inputs"].get("DCF") or {}).get("net_debt") or 100.0
    dr["financials_used"] = {**dr.get("financials_used", {}), "fx_rate": 1.0, "source_currency": "USD",
                             "net_debt_basis": {"balance_sheet_date": "2026-06-30", "components": {
                                 "feed_net_debt": nd_used + 50.0, "short_term_investments": 30.0, "long_term_investments": 20.0,
                                 "leases_removed": 0.0, "debt_like": 0.0, "result": nd_used}}}
    blob = build_workbook(run, "TEST")
    wb = load_workbook(io.BytesIO(blob))
    comps = [str(c.value) for c in wb["Comps"]["A"] if c.value]
    assert any("P/E (NTM) -- Forward P/E leg" in v for v in comps)
    summ = [str(c.value) for c in wb["Summary"]["A"] if c.value]
    assert any(v.startswith("Check: bridge") for v in summ)
    with tempfile.TemporaryDirectory() as td:
        p = os.path.join(td, "r3.xlsx")
        Path(p).write_bytes(blob)
        sol = {k.upper(): v for k, v in formulas.ExcelModel().loads(p).finish().calculate().items()}
        ws = wb["Summary"]
        r = next(c.row for c in ws["A"] if str(c.value or "").startswith("Check: bridge"))
        v = sol[f"'[R3.XLSX]SUMMARY'!B{r}".upper()]
        assert abs(getattr(v, "value", v)[0][0]) < 1e-6


def test_i14_the_double_count_guard_only_nets_a_pipeline_that_is_priced():
    # VRTX (2026-10-07): on the consolidated-DCF profile no rNPV add-on is priced, so the accepted pipeline must
    # not reduce the terminal replacement credit.
    import inspect
    src = inspect.getsource(d)
    assert 'if (profile_name or "") in _PIPELINE_ADDON_PROFILES else 0.0)' in src
    # Plan I2 (2026-10-07): big pharma now prices the pipeline add-on, so the guard applies there too
    assert "Big Pharma (Consolidated DCF)" in d._PIPELINE_ADDON_PROFILES
    from src.data import sector_profiles as sp_
    bp = sp_.INDUSTRY_VALUATION_PROFILES["Biopharma"]["Big Pharma (Consolidated DCF)"]["methods"]
    assert "rNPV (Pipeline)" in [m["name"] for m in bp] and abs(sum(m["weight"] for m in bp) - 1.0) < 1e-9


# ── Share-count intervention (owner, 2026-10-07, REGN) ─────────────────────────────────────────────────

def _failing_gf():
    rows = [{"year": 1, "net_income": 4.67e9}, {"year": 2, "net_income": 4.44e9}]
    return {"horizon_years": 2, "rows": rows, "deconstruction": {"eps_T_guided": 35.5},
            "invariants": [{"id": 3, "name": "Share-count integrity", "ok": False,
                            "detail": "guided EPS implies 125m shares in the target year vs 107m today"}]}


def _research_block():
    return {"fiscal_year_1": "FY2026", "fiscal_year_2": "FY2027", "confidence": "MEDIUM",
            "estimates": {"base": {"revenue_growth_fy1": 0.15, "revenue_growth_fy2": 0.12, "ebitda_margin_fy1": 0.31,
                                   "ebitda_margin_fy2": 0.31, "eps_fy1": 32.0, "eps_fy2": 35.5}},
            "medium_term_target": {"metric": "eps", "target_year": "FY2029", "mid": 50}}


def test_share_count_failure_with_street_coverage_rebuilds_the_forecast_on_the_street():
    fwd = {"analyst_count_eps": 15, "_fy1_fy2": {
        "eps": {"bear": (45.7, 47.9), "base": (55.94, 61.85), "bull": (61.4, 69.6)},
        "revenue": {"bear": (16.5e9, 17.4e9), "base": (17.17e9, 18.6e9), "bull": (17.8e9, 19.8e9)}}}
    blk = d._share_count_intervention(_research_block(), _failing_gf(), fwd,
                                      {"fy1_period": "2026-12-31", "fy2_period": "2027-12-31"},
                                      14.34e9, 107e6, 1.09, "last four quarters", "USD")
    sci = blk["_share_count_intervention"]
    assert sci["action"] == "rebuilt_on_street"
    assert sci["eps_model_reported"] == pytest.approx(4.44e9 / 107e6)
    b = blk["estimates"]["base"]
    assert b["eps_fy2"] == pytest.approx(61.85) and "ebitda_margin_fy2" not in b
    assert b["revenue_growth_fy1"] == pytest.approx(17.17e9 / 14.34e9 - 1)
    assert blk["_eps_basis_ratio"] == 1.09 and blk["medium_term_target"] == {}
    assert blk["estimates"]["bear"]["eps_fy2"] == pytest.approx(47.9)


def test_share_count_failure_without_street_coverage_rebases_to_the_model():
    fwd = {"analyst_count_eps": 2, "_fy1_fy2": {"eps": {"base": (55.94, 61.85)}, "revenue": {"base": (17.17e9, 18.6e9)}}}
    blk = d._share_count_intervention(_research_block(), _failing_gf(), fwd, {}, 14.34e9, 107e6, 1.0, "n/a", "USD")
    assert blk["_share_count_intervention"]["action"] == "rebased_to_model"
    assert "eps_fy2" not in blk["estimates"]["base"] and blk["medium_term_target"] == {}
    assert d._share_count_intervention(_research_block(), {**_failing_gf(), "invariants": [{"id": 3, "ok": True}]},
                                       fwd, {}, 14.34e9, 107e6, 1.0, "n/a", "USD") is None


def test_an_adjusted_eps_block_with_its_ratio_builds_the_same_forecast_as_the_reported_block():
    from src.agents.analysis import guidance_forecast as gfm
    series = [{"revenue": 34e9, "ebit": 1.6e9, "net_income": 1.1e9, "interest_expense": 110e6, "depreciation_and_amortization": 180e6,
               "capital_expenditure": -120e6, "change_in_working_capital": -200e6, "shares_outstanding": 58e6, "invested_capital": 7e9, "share_buyback": -500e6},
              {"revenue": 38e9, "ebit": 1.9e9, "net_income": 1.3e9, "interest_expense": 110e6, "depreciation_and_amortization": 190e6,
               "capital_expenditure": -140e6, "change_in_working_capital": -250e6, "shares_outstanding": 56e6, "invested_capital": 7.5e9, "share_buyback": -600e6}]
    rep = {"fiscal_year_1": "FY2026", "fiscal_year_2": "FY2027", "confidence": "MEDIUM",
           "estimates": {"base": {"revenue_growth_fy1": 0.05, "revenue_growth_fy2": 0.06, "eps_fy1": 25.0, "eps_fy2": 27.0}}}
    adj = {**rep, "estimates": {"base": {"revenue_growth_fy1": 0.05, "revenue_growth_fy2": 0.06, "eps_fy1": 30.0, "eps_fy2": 32.4}},
           "_eps_basis_ratio": 1.2}
    kw = dict(scenario="base", series=series, profile_name="Commercial Biotech", sector="Biopharma", wacc=0.08, tgr=0.025,
              shares=56e6, net_debt=0.0, spot=400.0)
    a, b = gfm.build_forecast(rep, **kw), gfm.build_forecast(adj, **kw)
    assert [round(r["net_income"]) for r in a["rows"]] == [round(r["net_income"]) for r in b["rows"]]
    i3 = next(i for i in b["invariants"] if i["id"] == 3)
    assert i3["ok"] is not False
    # the final payload states the share-count status
    p = d._guidance_forecast_payload({**b, "share_count_intervention": {"action": "rebuilt_on_street"}})
    assert p["share_count_status"] == "CORRECTED" and p["share_count_intervention"]["action"] == "rebuilt_on_street"
    assert d._guidance_forecast_payload(b)["share_count_status"] in ("PASS", "n/a")
    bad = {**b, "invariants": [{"id": 3, "ok": False}]}
    assert d._guidance_forecast_payload(bad)["share_count_status"] == "UNRESOLVED"


def test_pipeline_tab_lists_assets_phase_peak_pos_and_ties_out():
    # Owner, 2026-10-07: every drug-developer workbook carries a Pipeline tab; the risk-adjusted PV is a formula
    # (unrisked PV x PoS) and the totals tie to the engine.
    formulas = pytest.importorskip("formulas")
    run = _wbt._run()
    dr = run["data"]["dcf_range"]["TEST"]
    assets = [{"name": "Cemdisiran", "indication": "gMG", "phase": "filed", "launch_year": 2027, "peak_sales_usd": 2.0e9,
               "economic_share": None, "royalty_payable": 0.12, "base_phase_pos": 0.85, "ta_multiplier": 1.0,
               "effective_pos": 0.85, "years_to_launch": 0.5, "unrisked_pv": 3.0e9, "risk_adjusted_pv": 2.55e9,
               "ptrs_basis": "filed base rate"},
              {"name": "Approved X", "indication": "y", "phase": "approved", "launch_year": 2024, "peak_sales_usd": 1.0e9,
               "base_phase_pos": 1.0, "ta_multiplier": 1.0, "effective_pos": 1.0, "years_to_launch": 0.0,
               "unrisked_pv": 1.0e9, "risk_adjusted_pv": 1.0e9}]
    for s_, mult in (("bear", 0.75), ("base", 1.0), ("bull", 1.25)):
        dr[s_]["rnpv_audit"] = {"assets": assets, "pipeline_pv": 3.55e9, "shares_diluted": 100e6, "n_assets": 2,
                                "peak_scenario_multiplier": mult}
        dr[s_]["pipeline_addon"] = {"per_share": 25.5, "legs": ["DCF"]}
    dr["pipeline_input"] = {"status": "accepted", "as_of": "2026-10-07",
                            "excluded_assets": [{"name": "Factor XI (REGN7508)", "reason": "no single-asset peak"}],
                            "approved_portfolio": ["Dupixent"], "assets": [{"name": "Cemdisiran", "phase": "filed", "peak_period": "2031E", "peak_source": "Argus"}],
                            "primary_sources": [{"name": "Cemdisiran", "verdict": "corroborated"}]}
    blob = build_workbook(run, "TEST")
    wb = load_workbook(io.BytesIO(blob))
    assert "Pipeline" in wb.sheetnames
    ws = wb["Pipeline"]
    col_a = [str(c.value) for c in ws["A"] if c.value]
    assert "Cemdisiran" in col_a and any("REGN7508" in v for v in col_a)
    with tempfile.TemporaryDirectory() as td:
        p = os.path.join(td, "pipe.xlsx")
        Path(p).write_bytes(blob)
        sol = {k.upper(): v for k, v in formulas.ExcelModel().loads(p).finish().calculate().items()}
    checks = [c.row for c in ws["L"] if c.value == "Check"]
    assert len(checks) == 2                                   # pipeline PV and the add-on per share
    for r in checks:
        v = sol[f"'[PIPE.XLSX]PIPELINE'!M{r}".upper()]
        assert abs(getattr(v, "value", v)[0][0]) < 1e-6, r


def test_share_count_failure_on_the_streets_adjusted_basis_converts_instead_of_discarding():
    # BLK (2026-10-07): the research's as-adjusted EPS matches the street; it is a basis, not an error.
    fwd = {"analyst_count_eps": 18, "_fy1_fy2": {"eps": {"base": (34.0, 36.0)}, "revenue": {"base": (25e9, 27e9)}}}
    gfx = {**_failing_gf(), "deconstruction": {"eps_T_guided": 35.0}}
    blk = d._share_count_intervention(_research_block(), gfx, fwd, {}, 22e9, 161e6, 1.16, "last four quarters", "USD")
    assert blk["_share_count_intervention"]["action"] == "converted_to_reported"
    assert blk["_eps_basis_ratio"] == 1.16 and blk["estimates"]["base"]["eps_fy2"] == 35.5   # research estimates kept


def test_share_count_intervention_survives_null_fields_in_the_research_block():
    # REGN prod 2026-10-07: guidance / eps / estimates present but null -> TypeError on item assignment.
    fwd = {"analyst_count_eps": 15, "_fy1_fy2": {"eps": {"base": (55.94, 61.85)}, "revenue": {"base": (17.17e9, 18.6e9)}}}
    for blk0 in ({**_research_block(), "guidance": None}, {**_research_block(), "guidance": {"eps": None}},
                 {**_research_block(), "medium_term_target": None}):
        blk = d._share_count_intervention(blk0, _failing_gf(), fwd, {}, 14.34e9, 107e6, 1.09, "x", "USD")
        assert blk["_share_count_intervention"]["action"] == "rebuilt_on_street"
        assert blk["guidance"]["eps"]["currency"] == "USD"
    blk = d._share_count_intervention({**_research_block(), "estimates": None, "medium_term_target": None},
                                      _failing_gf(), {"analyst_count_eps": 1}, {}, 14.34e9, 107e6, 1.0, "x", "USD")
    assert blk["_share_count_intervention"]["action"] == "rebased_to_model"


def test_product_build_decomposes_the_forecast_by_product_and_costs_and_the_tab_ties():
    # Owner, 2026-10-07: biotech revenue by product (partnership revenue tagged) and a line-by-line cost build,
    # both a decomposition of the forecast the DCF runs on.
    formulas = pytest.importorskip("formulas")
    from src.agents.analysis import product_build as pbm
    from src.data import valuation_constants as vc
    entry = vc.load()["franchise_loe"]["entries"]["REGN"]
    sched = [0.20, 0.08, 0.04, 0.04, 0.03, 0.025, 0.025, 0.025, 0.025, 0.025]
    loe = d._franchise_loe_overlay(entry, 2025, sched, 0.0825, 0.025, profile_name="Commercial Biotech")
    rev, rows = 14.34e9, []
    for t, g in enumerate(loe["growth_schedule"], start=1):
        rev *= (1 + g)
        rows.append({"year": t, "revenue": rev, "ebit": rev * 0.33})
    series = [{"period": f"{y}-12-31", "revenue": r_, "cost_of_revenue": r_ * 0.15, "research_and_development": r_ * 0.38,
               "operating_expense": r_ * 0.58} for y, r_ in ((2023, 13.1e9), (2024, 14.2e9), (2025, 14.34e9))]
    pb = pbm.build(rows, series, 2025, loe=loe, entry=entry)
    names = {ln["name"]: ln for ln in pb["products"]}
    assert any("Dupixent" in k for k in names) and any("EYLEA HD" in k for k in names) and any("Libtayo" in k for k in names)
    dup = next(v for k, v in names.items() if "Dupixent" in k)
    assert dup["type"] == "partnership"                                   # Sanofi profit share
    for t in range(10):
        assert sum(ln["values"][t] for ln in pb["products"]) + pb["other"][t] == pytest.approx(rows[t]["revenue"])
    assert dup["values"][9] < dup["values"][4]                            # Dupixent erodes after its 2031 LOE
    cs = pb["costs"]
    assert cs["ratios"]["cost_of_revenue"] == pytest.approx(0.15) and cs["ratios"]["sga_and_other"] == pytest.approx(0.20)
    assert cs["ebit_build"][0] + cs["margin_path"][0] == pytest.approx(rows[0]["ebit"])
    run = _wbt._run()
    run["data"]["dcf_range"]["TEST"]["product_build"] = pb
    blob = build_workbook(run, "TEST")
    wb = load_workbook(io.BytesIO(blob))
    assert "Products" in wb.sheetnames
    ws = wb["Products"]
    with tempfile.TemporaryDirectory() as td:
        p = os.path.join(td, "prod.xlsx")
        Path(p).write_bytes(blob)
        sol = {k.upper(): v for k, v in formulas.ExcelModel().loads(p).finish().calculate().items()}
    checks = [c.row for c in ws["A"] if c.value == "Check"]
    assert len(checks) == 2
    for r in checks:
        for col in ("D", "H", "M"):
            v = sol[f"'[PROD.XLSX]PRODUCTS'!{col}{r}".upper()]
            assert abs(getattr(v, "value", v)[0][0]) < 1e-6, (r, col)


def test_a_drug_past_its_loe_never_grows_with_the_company():
    # REGN EYLEA 2 mg (2026-10-07): already eroding, it rose 3.3 -> 4.2bn on the company path once the curve ran out.
    entry = {"total_revenue": 100.0, "drugs": [{"name": "Eylea 2mg", "revenue_fy": 20.0, "loe_year": 2024, "modality": "biologic"},
                                               {"name": "Future", "revenue_fy": 30.0, "loe_year": 2030, "modality": "biologic"}]}
    o = d._franchise_loe_overlay(entry, 2025, [0.06] * 10, 0.08, 0.025)
    U = [1.06 ** t for t in range(1, 11)]
    for dp in o["drug_paths"]:
        vals = [U[t] * dp["path"][t] for t in range(10)]          # the drug's revenue relative to FY0 revenue
        start = 2 if dp["name"] == "Eylea 2mg" else 4            # after the covered years / from the year before LOE
        assert all(vals[t + 1] <= vals[t] + 1e-12 for t in range(start, 9)), dp["name"]


def test_e2_ev_ebitda_reverts_to_street_ebitda_when_the_model_is_on_another_basis():
    # AMGN (2026-10-07): reported EBITDA $16.3bn against $21.7bn consensus under an adjusted peer multiple.
    def sc(s_, cur, cons):
        return {"ebitda": {s_: cur}, "_consensus": {"ebitda": {s_: cons}}, "_source": {"ebitda": {s_: "guidance-derived FY+1 revenue x margin"}}}
    out, dec, flag = d._street_basis_guard(sc("base", 16.3e9, 21.7e9), "base", None, 20, "ebitda")
    assert dec is True and out["ebitda"]["base"] == 21.7e9 and "Forward EV/EBITDA" in flag and "21.70bn" in flag
    out_b, _, _ = d._street_basis_guard(sc("bull", 18e9, 23e9), "bull", dec, 20, "ebitda")
    assert out_b["ebitda"]["bull"] == 23e9
    _, dec2, _ = d._street_basis_guard(sc("base", 20.9e9, 21.7e9), "base", None, 20, "ebitda")
    assert dec2 is False


def test_e3_ddm_rolls_when_buybacks_are_a_material_share_of_the_cash_returned():
    import inspect
    src = inspect.getsource(d)
    assert d._DDM_MAX_BUYBACK_SHARE == 0.30
    assert "_bb_share >= _DDM_MAX_BUYBACK_SHARE" in src and "a dividend model prices only part of the payout" in src


def test_e4_agency_rating_replaces_the_synthetic_one_and_research_risk_is_waived_for_steady_earners():
    from src.data import sector_profiles as sp_
    assert sp_.agency_rating_bucket("BBB+") == "BBB" and sp_.agency_rating_bucket("A3") == "A"
    syn = sp_.get_cost_of_debt(interest_coverage=30.0, sector="Biopharma")
    agy = sp_.get_cost_of_debt(interest_coverage=30.0, sector="Biopharma", rating_override="BBB+")
    assert syn["rating"] == "AAA" and agy["rating"] == "BBB" and agy["cost_of_debt"] > syn["cost_of_debt"]
    import inspect
    src = inspect.getsource(d)
    assert 'agency_rating=_agency_rating(ticker)' in src and 'is_biopharma_sector(sector)' in src and '(_prof3 or _pipe_priced)' in src
    assert d._agency_rating("GILD") is None or isinstance(d._agency_rating("GILD"), str)


def test_i4_an_already_eroding_product_is_shown_eroding_from_the_base_year():
    entry = {"total_revenue": 100.0, "drugs": [{"name": "Prolia", "revenue_fy": 12.0, "loe_year": 2025, "modality": "biologic"}]}
    o = d._franchise_loe_overlay(entry, 2025, [0.05] * 10, 0.08, 0.025)
    U = [1.05 ** t for t in range(1, 11)]
    vals = [U[t] * o["drug_paths"][0]["path"][t] * 100.0 for t in range(10)]
    assert vals[0] < 12.0 and vals[1] < vals[0]                    # eroding in the guided years too


def test_e5_leg_dispersion_is_measured_on_the_weighted_legs():
    w = [{"method": "DCF", "value_key": "DCF", "weight": 0.5}, {"method": "Forward P/E", "value_key": "Forward P/E", "weight": 0.3},
         {"method": "SOTP", "value_key": "SOTP", "weight": 0.0}]
    disp = d._leg_dispersion(w, {"DCF": 121.0, "Forward P/E": 386.0, "SOTP": 10.0})
    assert disp["ratio"] == pytest.approx(386.0 / 121.0, abs=1e-3) and disp["lo_leg"] == "DCF"   # the unweighted leg ignored
    assert disp["ratio"] > d._LEG_DISPERSION_MAX
    assert d._leg_dispersion(w[:1], {"DCF": 1.0}) is None
    import inspect
    assert 'DEGRADED: the forecast\'s share-count check is unresolved' in inspect.getsource(d) or "share-count check is unresolved" in inspect.getsource(d)


def test_e6_terminal_check_uses_the_forward_peer_median():
    assert d._peer_fwd_ev_ebitda({"ev_ebitda": 24.9, "ev_ebitda_ntm": 12.4}) == 12.4
    assert d._peer_fwd_ev_ebitda({"ev_ebitda": 15.0}) == 15.0 and d._peer_fwd_ev_ebitda(None) is None


def test_e7_growth_premium_is_capped_when_the_forecast_does_not_outgrow_peers():
    fwd = {"_fy1_fy2": {"revenue": {"base": (100.0, 110.0)}}}
    peer = {"growth_avg": 0.04}
    f_raw, _ = d._forward_growth_adjustment(fwd, peer)
    f_cap, note = d._forward_growth_adjustment(fwd, peer, forecast_cagr=0.01)
    assert f_raw > 1.0 and f_cap == 1.0 and "capped" in note
    assert d._forward_growth_adjustment(fwd, peer, forecast_cagr=0.06)[0] == f_raw


def test_e8_a_net_debt_jump_since_the_year_end_is_flagged_as_a_possibly_unvalued_acquisition():
    import inspect
    src = inspect.getsource(d)
    assert "UNVALUED ACQUISITION?" in src and "_drop >= 0.15 * _rev_src" in src


def test_e9_guidance_that_is_not_company_level_is_flagged():
    est = {"guidance": {"revenue": {"low": 1, "mid": 1, "high": 1, "scale": "bn", "currency": "USD"}},
           "estimates": {"base": {"revenue_growth_fy1": 0.095}}}
    fwd = {"analyst_count_revenue": 20, "_fy1_fy2": {"revenue": {"base": (29.6e9, 31.0e9)}}}
    chk = d._guidance_scope_checks(est, 29.4e9, fwd, 29.4e9)
    names = [c["check"] for c in chk]
    assert "guided revenue level" in names and "research growth vs street" in names
    ok = d._guidance_scope_checks({"guidance": {"revenue": {"mid": 29.5, "scale": "bn"}},
                                   "estimates": {"base": {"revenue_growth_fy1": 0.02}}}, 29.4e9, fwd, 29.4e9)
    assert ok == []
    p = d._guidance_estimates_payload(est, None, chk)
    assert p["scope_checks"] == chk


def test_e10_reverse_dcf_solves_the_growth_shift_that_meets_the_price():
    kw = dict(revenue_base=1000.0, fcf_margin_base=0.2, growth_rate=0.05, margin_delta_per_year=0.0, wacc=0.09, tgr=0.025,
              fcf_floor=-0.05, net_debt=100.0, shares=10.0, growth_schedule=[0.05] * 10)
    iv = d._project_dcf(**kw)[0]
    rd = d._reverse_dcf(kw, iv * 1.3)
    assert rd["solved"] and rd["growth_shift"] > 0 and rd["implied_cagr10"] > rd["model_cagr10"]
    shifted = dict(kw, growth_schedule=[0.05 + rd["growth_shift"]] * 10)
    assert d._project_dcf(**shifted)[0] == pytest.approx(iv * 1.3, rel=1e-4)
    assert d._reverse_dcf(kw, iv)["growth_shift"] == pytest.approx(0.0, abs=1e-4)
