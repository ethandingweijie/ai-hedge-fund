"""Report families (owner, 2026-09-27): every profile renders under one family, the family's
key-financials rows read the per-year raw financials, and both renderers use them."""
import io

import pytest
from openpyxl import load_workbook

from src.data import report_families as rf
from src.data import sector_profiles as sp


def _all_profiles():
    return {p for profs in sp.INDUSTRY_VALUATION_PROFILES.values() for p in profs}


def test_every_profile_maps_to_exactly_one_family_and_no_family_lists_a_ghost():
    listed = [p for f, spec in rf.REPORT_FAMILIES.items() for p in spec["profiles"]]
    assert len(listed) == len(set(listed)), "a profile is listed under two families"
    unmapped = sorted(p for p in _all_profiles() if rf.report_family_for(p) == "Operating company")
    assert unmapped == [], unmapped
    ghosts = sorted(set(listed) - _all_profiles())
    assert ghosts == [], ghosts
    assert rf.report_family_for("Money Center Bank") == "Banks"
    assert rf.report_family_for("Card Issuer & Consumer Lender") == "Banks"
    assert rf.report_family_for("Insurance (P&C)") == "Insurance"
    assert rf.report_family_for("Upstream Oil & Gas") == "Energy and resources"
    assert rf.report_family_for("Mature SaaS") == "Technology, telecom and media"
    assert rf.report_family_for("nonsense") == "Operating company"


def test_every_family_row_resolves_to_a_known_spec():
    for fam, spec in rf.REPORT_FAMILIES.items():
        assert spec["rows"], fam
        for k in spec["rows"]:
            assert k in rf._R, (fam, k)
        for f in ("exposition", "skeleton", "signals"):
            assert isinstance(spec.get(f), list), (fam, f)


RAW = {"currency": "USD",
       "FY2024": {"revenue": 100e9, "net_income": 10e9, "ebit": 15e9, "ebitda": 20e9, "free_cash_flow": 12e9,
                  "net_debt": 30e9, "total_equity": 50e9, "shares_outstanding": 1e9, "dividends_per_share": 2.0,
                  "capital_expenditure": -8e9, "stock_based_compensation": 3e9},
       "FY2025": {"revenue": 110e9, "net_income": 11e9, "ebit": 16.5e9, "ebitda": 22e9, "free_cash_flow": 13e9,
                  "net_debt": 25e9, "total_equity": 55e9, "shares_outstanding": 1e9, "book_value_per_share": 55.5,
                  "dividends_per_share": 2.2, "capital_expenditure": -9e9, "stock_based_compensation": 3.3e9}}


def test_family_rows_read_and_derive_from_the_raw_financials():
    fys, rows = rf.family_rows(RAW, "Banks")
    assert fys == ["FY2024", "FY2025"]
    d = {label: vals for label, _, vals in rows}
    assert d["Book value / share"] == [pytest.approx(50.0), 55.5]           # derived, then the stated figure
    assert d["Return on equity"] == [pytest.approx(0.2), pytest.approx(0.2)]
    assert "FCF" not in d and "Net debt / (cash)" not in d
    fys, rows = rf.family_rows(RAW, "Technology, telecom and media")
    d = {label: vals for label, _, vals in rows}
    assert d["SBC / revenue"] == [pytest.approx(0.03), pytest.approx(0.03)]
    assert d["Capex / revenue"] == [pytest.approx(0.08), pytest.approx(0.0818, abs=1e-3)]   # absolute of the outflow
    fys, rows = rf.family_rows(RAW, "Energy and resources")
    d = {label: vals for label, _, vals in rows}
    assert d["Net debt / EBITDA"] == [pytest.approx(1.5), pytest.approx(25 / 22)]
    assert rf.family_rows(None, "Banks") == ([], []) and rf.family_rows({"currency": "USD"}, "Banks") == ([], [])
    assert rf.format_value(12e9, "bn") == "12.0B" and rf.format_value(0.2, "pct") == "20.0%" and rf.format_value(None, "x") == "n/a"


def test_the_workbook_summary_carries_the_family_block():
    import sys
    sys.path.insert(0, "tests")
    from test_valuation_workbook import _run, _statements
    from src.utils.valuation_workbook import build_workbook
    run = _run()
    run["data"]["raw_financials"] = RAW
    run["data"]["dcf_range"]["TEST"]["profile"] = "Money Center Bank"
    wb = load_workbook(io.BytesIO(build_workbook(run, "TEST", load_statements=_statements)))
    ws = wb["Summary"]
    cells = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
    assert "Key metrics (Banks)" in cells and "Book value / share" in cells and "Return on equity" in cells



def test_the_workbook_carries_a_family_tab_with_the_checklist():
    import sys
    sys.path.insert(0, "tests")
    from test_valuation_workbook import _run, _statements
    from src.utils.valuation_workbook import build_workbook
    run = _run()
    run["data"]["dcf_range"]["TEST"]["profile"] = "Integrated Oil & Gas"
    wb = load_workbook(io.BytesIO(build_workbook(run, "TEST", load_statements=_statements)))
    assert "Family" in wb.sheetnames
    cells = " | ".join(str(c.value) for row in wb["Family"].iter_rows() for c in row if c.value is not None)
    assert "Energy and resources" in cells and "Thesis skeleton" in cells and "Checklist from this run" in cells
    assert "mid-cycle multiple" in cells
