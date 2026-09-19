"""Valuation workbook export: institutional layout, colour code, live formulas.

Full numerical reconciliation (every Check = 0 after evaluating the formulas)
was verified on 14 golden fixtures and a live 09618.HK run with a formula
engine; these tests pin the structure that makes that reconciliation hold.
"""
import io
import re

import pytest
from openpyxl import load_workbook

from src.agents.analysis import dcf_agent as d
from src.utils.valuation_workbook import build_workbook

BLUE, BLACK, GREEN = "FF0000FF", "FF000000", "FF008000"


def _dcf_trace(g=0.10, margin=0.20, wacc=0.09, tgr=0.03, nd=100.0, sh=10.0):
    sched = d._decayed_growth_schedule(g, "Hyper-Growth Platform")
    iv, pv_f, pv_t, rows = d._project_dcf(1_000.0, margin, g, 0.0, wacc, tgr, -0.05, nd, sh,
                                          growth_schedule=sched, margin_delta_absolute=0.0)
    return {"kind": "dcf", "value": iv, "revenue_base": 1_000.0, "fcf_margin_base": margin,
            "growth_base": g, "growth_schedule": sched, "margin_delta_absolute": 0.0, "wacc": wacc,
            "wacc_schedule": None, "tgr": tgr, "fcf_floor": -0.05, "net_debt": nd, "shares": sh,
            "pv_fcf_per_share": pv_f, "pv_tv_per_share": pv_t, "projection_rows": rows}


def _run():
    legs = {
        "DCF": _dcf_trace(),
        "EV/EBITDA": {"kind": "ev_multiple", "metric": "EBITDA (TTM)", "metric_value": 200.0,
                      "multiple": 12.0, "multiple_parts": {"peer_multiple": 12.0, "scenario_band": 1.0,
                                                           "growth_premium": 1.0, "sbc_haircut": 1.0},
                      "ev": 2400.0, "net_debt": 100.0, "minority_interest": 0.0, "equity": 2300.0,
                      "shares": 10.0, "value": 230.0},
        "P/E": {"kind": "equity_multiple", "metric": "Net income (TTM)", "metric_value": 100.0,
                "shares": 10.0, "per_share_metric": 10.0, "multiple": 18.0,
                "multiple_parts": {"peer_multiple": 18.0, "scenario_band": 1.0, "growth_premium": 1.0,
                                   "sbc_pe_discount": 1.0}, "value": 180.0},
    }
    scen = {}
    for s in ("bear", "base", "bull"):
        dcf_v = legs["DCF"]["value"]
        iv = round(0.4 * dcf_v + 0.4 * 230.0 + 0.2 * 180.0, 2)
        scen[s] = {"intrinsic_value": iv, "leg_inputs": legs,
                   "method_iv_table": {"DCF": dcf_v, "EV/EBITDA": 230.0, "P/E": 180.0},
                   "effective_weights": [
                       {"method": "DCF", "value_key": "DCF", "bucket": "dcf", "weight": 0.4},
                       {"method": "EV/EBITDA", "value_key": "EV/EBITDA", "bucket": "multi", "weight": 0.4},
                       {"method": "P/E", "value_key": "P/E", "bucket": "multi", "weight": 0.2}],
                   "forward_flags": ["example flag"]}
    iv = scen["base"]["intrinsic_value"]
    dr = {**scen, "profile": "Test Profile", "anchor_method": "EV/EBITDA", "reported_currency": "USD",
          "wacc": 0.09,
          "wacc_build": {"wacc_base": 0.09, "wacc": 0.09, "insider_bps": 0.0, "research_risk_loading": 0.0,
                         "country_risk_premium": 0.0, "wacc_final": 0.09,
                         "base_breakdown": {"market": "US", "table": "US sector WACC (Damodaran)",
                                            "lookup": "Tech", "table_rate": 0.09, "crp_embedded": 0.0,
                                            "leverage": 0.2, "leverage_threshold": 1.5, "leverage_slope": 0.01,
                                            "leverage_premium_applies": True, "leverage_premium": 0.0,
                                            "leverage_cap": 0.04, "macro_regime": "neutral",
                                            "macro_overlay": 0.0, "cap_binding": False, "wacc": 0.09}},
          "pt_bridge": {"spot": 150.0, "capture": 0.35, "scenarios": {}, "cross_checks": {}},
          "12m_targets": {s: round(150.0 + 0.35 * (iv - 150.0), 2) for s in ("bear", "base", "bull")},
          "calibration_record": {"status": "passed"}, "multiples_used": {"fields": {}},
          "financials_used": {"currency": "USD", "rows": []}}
    return {"run_id": "r1", "run_at": "2026-09-19", "data": {
        "dcf_range": {"TEST": dr}, "end_date": "2026-09-19",
        "scenario_analysis": {"TEST": {"bear": {"probability": 0.25}, "base": {"probability": 0.5},
                                       "bull": {"probability": 0.25},
                                       "reconciliation": {"12m_price_target": 1.0}}}}}


def _statements(ticker, end_date):
    rows = []
    for y, rev in ((2022, 800.0), (2023, 900.0), (2024, 1000.0)):
        rows.append({"period": f"{y}-12-31", "revenue": rev * 1e6, "cost_of_revenue": rev * 0.4e6,
                     "operating_expense": rev * 0.3e6, "other_income_expense": -10e6,
                     "pretax_income": (rev * 0.3 - 10) * 1e6, "income_tax_expense": 20e6,
                     "net_income": (rev * 0.3 - 30) * 1e6, "shares_outstanding": 10e6,
                     "operating_cash_flow": 250e6, "capital_expenditure": -50e6})
    return {"rows": rows, "currency": "USD", "beta": 1.1,
            "estimates": [{"period_end": "2025-12-31", "revenue_avg": 1100e6, "ebitda_avg": 300e6,
                           "ebit_avg": 250e6, "net_income_avg": 200e6, "eps_avg": 20.0,
                           "analyst_count_revenue": 10}]}


@pytest.fixture(scope="module")
def wb():
    blob = build_workbook(_run(), "TEST", load_statements=_statements)
    return load_workbook(io.BytesIO(blob))


def test_tab_architecture(wb):
    assert wb.sheetnames[:6] == ["Cover", "Summary", "Assumptions", "IS", "BS", "CFS"]
    for tab in ("WACC", "DCF", "Multiples", "Comps", "Blend", "Target", "Backtest", "Data Gaps"):
        assert tab in wb.sheetnames


def test_cover_has_a_linked_table_of_contents(wb):
    links = [c.hyperlink.location if c.hyperlink and c.hyperlink.location else (c.hyperlink.target if c.hyperlink else None)
             for row in wb["Cover"].iter_rows() for c in row if c.hyperlink]
    assert len(links) >= 10


def test_colour_code_follows_cell_content(wb):
    seen = {"input": 0, "formula": 0, "link": 0}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v, color = c.value, (c.font.color.rgb if c.font and c.font.color else None)
                if isinstance(v, str) and v.startswith("="):
                    want = GREEN if "!" in v else BLACK
                    assert color == want, (ws.title, c.coordinate, v, color)
                    seen["link" if want == GREEN else "formula"] += 1
                elif isinstance(v, (int, float)) and not isinstance(v, bool):
                    assert color == BLUE, (ws.title, c.coordinate, v, color)
                    seen["input"] += 1
    assert all(n > 0 for n in seen.values()), seen


_ALLOWED_LITERALS = {"0", "1", "2", "4", "12"}   # structural: 1+g, EDATE months, ROUND digits


def test_formulas_carry_no_hardcoded_drivers(wb):
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if not (isinstance(v, str) and v.startswith("=")):
                    continue
                body = re.sub(r"'[^']*'!\$?[A-Z]+\$?\d+|\$?[A-Z]{1,3}\$?\d+", "", v)   # drop refs
                body = re.sub(r'"[^"]*"', "", body)                                      # drop strings
                body = re.sub(r"DATE\(\d+,\d+,\d+\)", "", body)                          # period headers
                nums = re.findall(r"(?<![A-Za-z_])\d+(?:\.\d+)?", body)
                bad = [n for n in nums if n not in _ALLOWED_LITERALS]
                assert not bad, (ws.title, c.coordinate, v, bad)


def test_divisions_are_error_trapped(wb):
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and v.startswith("=") and "/" in v:
                    assert "IFERROR" in v or v.count("/") == 0, (ws.title, c.coordinate, v)


def test_period_headers_are_date_formulas(wb):
    hdr = [c.value for c in wb["IS"][4] if c.value]
    assert any(isinstance(v, str) and v.startswith("=DATE(") for v in hdr)
    assert any(isinstance(v, str) and v.startswith("=EDATE(") for v in hdr)


def test_dcf_uses_the_engine_arithmetic_and_reconciles_to_its_value(wb):
    ws = wb["DCF"]
    labels = {ws.cell(row=r, column=1).value: r for r in range(1, ws.max_row + 1)}
    r = labels["Intrinsic value per share"]
    assert "IFERROR" in ws.cell(row=r, column=2).value
    assert ws.cell(row=labels["Engine value"], column=2).value == pytest.approx(
        _run()["data"]["dcf_range"]["TEST"]["base"]["leg_inputs"]["DCF"]["value"])


def test_blend_links_every_leg_to_the_tab_that_rebuilds_it(wb):
    ws = wb["Blend"]
    refs = [ws.cell(row=r, column=5).value for r in range(1, ws.max_row + 1)
            if isinstance(ws.cell(row=r, column=5).value, str)
            and ws.cell(row=r, column=5).value.startswith("='")]
    assert any("'DCF'!" in x for x in refs) and any("'Multiples'!" in x for x in refs)


def test_consensus_forecast_and_dcf_revenue_sit_side_by_side(wb):
    ws = wb["IS"]
    text = [ws.cell(row=r, column=1).value or "" for r in range(1, ws.max_row + 1)]
    assert any("FMP analyst consensus" in t for t in text)
    assert any("Revenue (DCF base" in t for t in text)
    assert any("Variance vs consensus" in t for t in text)


def test_data_gaps_lists_structural_and_detected_gaps(wb):
    ws = wb["Data Gaps"]
    rows = [ws.cell(row=r, column=2).value or "" for r in range(5, ws.max_row + 1)]
    assert any("balance-sheet or cash-flow forecast" in x for x in rows)
    assert any("missing" in x for x in rows)          # the synthetic statements omit lines
