"""Steps two and three (owner, 2026-10-03): family-specific guided metrics in the research extractor,
and the bank / insurer earnings-and-capital model whose RoTE, FY+1 book and EPS price the legs."""
import copy
import importlib.util
import io
from pathlib import Path

import pytest
from openpyxl import load_workbook

_wspec = importlib.util.spec_from_file_location("_wbt", Path(__file__).resolve().parent / "test_valuation_workbook.py")
_wbt = importlib.util.module_from_spec(_wspec)
_wspec.loader.exec_module(_wbt)

from src.agents.analysis import bank_model as bmod
from src.agents.industry import deep_research as dr

MR = {"total_equity": 68.9e9, "total_assets": 897.5e9, "net_income": 11.0e9, "shares_outstanding": 2.85e9, "goodwill": 6.3e9, "intangible_assets": 0.0,
      "interest_income": 40e9, "interest_expense": 25e9, "pretax_income": 12.6e9, "income_tax_expense": 1.6e9, "report_period": "2025-12-31"}
BM = {"cet1_ratio": 0.146, "nim_pct": 0.0187, "efficiency_ratio": 0.39, "credit_cost_bps": 18.5, "dividend_payout_ratio": 0.375, "loan_growth_yoy": 0.06,
      "management_target_roe": 0.165, "cost_of_equity": 0.10}
BLOCK = {"fiscal_year_1": "FY2026", "confidence": "HIGH", "estimates": {"base": {"revenue_growth_fy1": 0.022, "revenue_growth_fy2": 0.035}},
         "family_metrics": {"family": "bank", "bank": {"loan_growth_fy1": 0.05, "nim_fy1": 0.0205, "cost_to_income_fy1": 0.40, "credit_cost_bps_fy1": 20.0, "payout_ratio": 0.55, "cet1_target": 0.135}}}


@pytest.fixture(autouse=True)
def _live_clock_after():
    """The D05.SI golden below replays IN-PROCESS, and replay pins the clock to the fixture's capture
    date for every first-party module. Hand the interpreter back live: left frozen, every later test in
    the session ran on 2026-09 (the comps slot gate, the fast-path reuse windows, the freshness pulse
    date and the research-age tests all failed only in a full run)."""
    yield
    from src.memory.golden_replay import unfreeze_clock
    unfreeze_clock()


def _opening():
    return bmod.opening_from_line_items(MR, 2.85e9, 22.3e9, "FY2025")


def test_the_extractor_schema_and_normaliser_carry_the_family_metrics():
    assert '"family_metrics"' in dr._GUIDANCE_ESTIMATES_SYSTEM and "cost-to-income" in dr._GUIDANCE_ESTIMATES_SYSTEM
    out = dr._normalize_guidance_estimates({"estimates": {"base": {"revenue_growth_fy1": 0.03}}, "confidence": "HIGH",
                                            "family_metrics": {"family": "bank", "bank": {"loan_growth_fy1": "5%", "nim_fy1": 0.021, "credit_cost_bps_fy1": "20", "cet1_target": 0.135,
                                                                                            "roe_target": {"value": "17%", "target_year": "FY2027"}}, "insurer": None}})
    fm = out["family_metrics"]
    assert fm["family"] == "bank" and fm["bank"]["loan_growth_fy1"] == 0.05 and fm["bank"]["credit_cost_bps_fy1"] == 20.0 and fm["bank"]["roe_target"] == {"value": 0.17, "target_year": "FY2027"}
    assert fm["insurer"] is None
    assert dr._normalize_guidance_estimates({"estimates": {"base": {"revenue_growth_fy1": 0.03}}, "family_metrics": {"family": None, "bank": None, "insurer": None}})["family_metrics"] is None


def test_assumptions_prefer_guidance_then_the_extraction_then_the_filings_then_defaults():
    a = bmod.bank_assumptions(_opening(), BM, BLOCK, coe=0.086, target_roe=0.165, cet1_target=0.14)
    assert a["asset_growth_fy1"] == {"value": 0.05, "source": "guidance: loan growth FY+1", "needed_for": "earning assets and loans, year 1"}
    assert a["asset_growth_fy2"]["source"] == "guidance: loan growth FY+1 carried"
    assert a["nim_fy1"]["value"] == 0.0205 and a["nim_fy1"]["source"].startswith("guidance") and a["nim_fy2"]["source"].endswith("(held)")
    assert a["cost_to_income_fy1"]["value"] == 0.40 and a["credit_cost_bps_fy1"]["value"] == 20.0 and a["payout_ratio"]["value"] == 0.55 and a["cet1_target"]["value"] == 0.135
    assert a["fee_income_growth_fy1"]["source"].startswith("estimate") and a["tax_rate"]["source"].startswith("line items") and a["tax_rate"]["value"] == pytest.approx(1.6 / 12.6)
    assert a["loan_share_of_assets"]["source"] == "default (PROPOSED)" and a["coe"]["value"] == 0.086
    assert set(a["_guided_fields"]["value"]) == {"loan_growth_fy1", "nim_fy1", "cost_to_income_fy1", "credit_cost_bps_fy1", "payout_ratio", "cet1_target"}
    b = bmod.bank_assumptions(_opening(), BM, {"estimates": {"base": {"revenue_growth_fy1": 0.022}}}, coe=0.086)
    assert b["nim_fy1"]["source"] == "bank metrics: NIM" and b["cet1_target"]["source"] == "default (PROPOSED)" and b["_guided_fields"]["value"] == []
    c = bmod.bank_assumptions(_opening(), None, None)
    assert c["nim_fy1"]["source"].startswith("line items") and c["nim_fy1"]["value"] == pytest.approx(15e9 / (897.5e9 * 0.9))


def test_the_bank_model_ties_balances_capital_and_per_share_lines_and_passes_its_suite():
    op = _opening()
    a = bmod.bank_assumptions(op, BM, BLOCK, coe=0.086, target_roe=0.165, cet1_target=0.14)
    m = bmod.build_bank(op, a, scenario="base")
    r = m["rows"]
    assert m["kind"] == "bank" and m["fy_labels"] == ["FY2026E", "FY2027E", "FY2028E", "FY2029E", "FY2030E"] and m["reconciliation"]["ok"]
    # NII = average earning assets × NIM; income = NII + fees; opex = cost-to-income × income; provisions = credit cost × loans
    assert r["net_interest_income"][0] == pytest.approx(r["earning_assets_avg"][0] * 0.0205)
    assert r["total_income"][0] == pytest.approx(r["net_interest_income"][0] + r["fee_income"][0])
    assert -r["operating_expenses"][0] == pytest.approx(0.40 * r["total_income"][0]) and -r["provisions"][0] == pytest.approx(r["loans_avg"][0] * 20.0 / 1e4)
    assert r["net_income"][0] == pytest.approx(r["pre_provision_profit"][0] + r["provisions"][0] + r["tax"][0])
    # capital: equity rolls with retained earnings; dividends at the guided payout; CET1 above target
    assert r["equity"][0] == pytest.approx(op["equity"] + r["net_income"][0] + r["dividends"][0] + r["buybacks"][0])
    assert -r["dividends"][0] == pytest.approx(0.55 * r["net_income"][0]) and r["cet1_ratio"][0] >= 0.135
    assert r["bvps"][0] == pytest.approx(r["equity"][0] / r["shares"][0]) and r["eps"][0] == pytest.approx(r["net_income"][0] / r["shares"][0])
    assert m["bvps_fy1"] == r["bvps"][0] and m["eps_fy1"] == r["eps"][0] and 0.05 < m["steady_rote"] < 0.40
    assert set(m["guided_fields"]) == {"loan_growth_fy1", "nim_fy1", "cost_to_income_fy1", "credit_cost_bps_fy1", "payout_ratio", "cet1_target"}
    # DBS-sized earnings on DBS-sized inputs: around S$11-13bn, not the inflated figure a gross-interest "revenue" would give
    assert 9e9 < r["net_income"][0] < 15e9
    bear, bull = bmod.build_bank(op, a, scenario="bear"), bmod.build_bank(op, a, scenario="bull")
    assert bear["rows"]["net_income"][0] < r["net_income"][0] < bull["rows"]["net_income"][0]
    assert len(bmod.coverage(m)) > 10 and any(c["status"] == "from management guidance" for c in bmod.coverage(m))


def test_a_capital_shortfall_cuts_the_dividend_before_the_ratio_breaches():
    op = _opening()
    a = bmod.bank_assumptions(op, {**BM, "cet1_ratio": 0.12}, {**BLOCK, "family_metrics": {"family": "bank", "bank": {"payout_ratio": 1.0, "cet1_target": 0.16, "loan_growth_fy1": 0.15, "loan_growth_fy2": 0.15}}}, coe=0.086)
    m = bmod.build_bank(op, a, scenario="base")
    assert m["reconciliation"]["ok"] and any(c["distribution_cut"] > 0 for c in m["checks"]) and any("cut" in n for n in m["notes"])
    # while the ratio sits below the target nothing is distributed; once it clears, dividends resume
    for c, div in zip(m["checks"], m["rows"]["dividends"]):
        if c["cet1_ratio"] < 0.16 - 1e-9:
            assert div == pytest.approx(0.0)
    assert any("capital shortfall" in n for n in m["notes"]) or m["checks"][-1]["cet1_ratio"] >= 0.16 - 1e-9


def test_the_insurer_variant_runs_on_premiums_combined_ratio_and_float():
    op = bmod.opening_from_line_items({"total_equity": 20e9, "total_assets": 120e9, "net_income": 2.4e9, "shares_outstanding": 1e9, "goodwill": 1e9, "pretax_income": 3.0e9, "income_tax_expense": 0.6e9,
                                       "report_period": "2025-12-31"}, 1e9, 30e9, "FY2025")
    a = bmod.insurer_assumptions(op, {"combined_ratio": 0.95, "investment_yield": 0.04}, {"family_metrics": {"family": "insurer", "insurer": {"premium_growth_fy1": 0.06, "combined_ratio_fy1": 0.93}}}, coe=0.09)
    m = bmod.build_insurer(op, a, scenario="base")
    r = m["rows"]
    assert m["kind"] == "insurer" and m["reconciliation"]["ok"] and a["combined_ratio_fy1"]["source"].startswith("guidance") and a["investment_yield_fy1"]["source"].startswith("insurance metrics")
    assert r["underwriting_result"][0] == pytest.approx(r["net_earned_premiums"][0] * (1 - 0.93)) and r["investment_income"][0] == pytest.approx(r["float_avg"][0] * 0.04)
    assert r["equity"][0] == pytest.approx(op["equity"] + r["net_income"][0] + r["dividends"][0]) and m["bvps_fy1"] == pytest.approx(r["equity"][0] / 1e9)


def _bank_payload():
    """A stored bank run: GGM (P/B) at 45%, forward P/E at 15%, with the model's context."""
    op = _opening()
    a = bmod.bank_assumptions(op, BM, BLOCK, coe=0.086, target_roe=0.165, cet1_target=0.14)
    m = bmod.build_bank(op, a, scenario="base")
    m["coverage"] = bmod.coverage(m)
    ggm = {"kind": "ggm", "target_pb": 2.2, "assumptions": {"roe": 0.165, "coe": 0.086, "g": 0.033, "roe_book": 0.15, "bvps": 24.28, "tbv_per_share": 22.06}, "value_before_band": 53.4, "scenario_band": 1.0, "value": 53.4}
    fpe = {"kind": "equity_multiple", "metric": "EPS (NTM consensus, base)", "metric_value": 4.1, "per_share_metric": 4.1, "multiple": 10.2, "value": 41.8}
    legs = {"GGM (P/B)": ggm, "Forward P/E": fpe, "Residual Income": {"value": 36.8}, "P/E (norm)": {"kind": "equity_multiple", "metric": "Net income (5y normalised)", "metric_value": 12.7e9, "multiple": 10.4, "value": 46.8}}
    eff = [{"method": "GGM (P/B)", "value_key": "GGM (P/B)", "weight": 0.45}, {"method": "Residual Income", "value_key": "Residual Income", "weight": 0.25},
           {"method": "Forward P/E", "value_key": "Forward P/E", "weight": 0.15}, {"method": "P/E (norm)", "value_key": "P/E (norm)", "weight": 0.15}]
    iv = 0.45 * 53.4 + 0.25 * 36.8 + 0.15 * 41.8 + 0.15 * 46.8
    spot, cap = 77.2, 0.5
    scen = {"intrinsic_value": round(iv, 2), "leg_inputs": copy.deepcopy(legs), "effective_weights": eff,
            "method_iv_table": {"GGM (P/B)": 53.4, "Residual Income": 36.8, "Forward P/E": 41.8, "P/E (norm)": 46.8}, "forward_flags": []}
    tgt = spot + cap * (iv - spot)
    dr_ = {"profile": "Money Center Bank (SG)", "reported_currency": "SGD", "base": scen, "bear": copy.deepcopy(scen), "bull": copy.deepcopy(scen), "shares_outstanding": 2.85e9,
           "pt_bridge": {"spot": spot, "capture": cap, "scenarios": {s: {"intrinsic_value": round(iv, 2), "target": round(tgt, 2)} for s in ("bear", "base", "bull")}},
           "12m_targets": {s: round(tgt, 2) for s in ("bear", "base", "bull")}, "guidance_estimates": {**BLOCK, "applied": False},
           "three_statements": m, "forecast_context": {"history": {"revenue": 22.3e9}, "inputs": {"spot": spot}, "statements_family_ok": False,
                                                       "bank_model": {"kind": "bank", "opening": op, "assumptions": a}}}
    sa = {"bear": {"probability": 0.3}, "base": {"probability": 0.5}, "bull": {"probability": 0.2}, "12m_price_target": round(tgt, 2), "expected_value": round(iv, 2), "current_price": spot,
          "reconciliation": {"current_price": spot, "blended_iv": round(iv, 2), "expected_value": round(iv, 2), "12m_price_target": round(tgt, 2)}}
    data = {"dcf_range": {"D05.SI": dr_}, "scenario_analysis": {"D05.SI": sa}, "decisions": {"D05.SI": {"action": "SELL", "price_target": round(tgt, 2), "rationale": "x"}}, "sector": "Financials"}
    return {"run_id": "run-b", "ticker": "D05.SI", "run_at": "2026-10-03T00:00:00", "data": data, "decisions": copy.deepcopy(data["decisions"])}


def test_the_workbench_rebuilds_the_bank_model_and_reprices_the_ggm_and_forward_pe_legs():
    from app.backend.services import estimate_override_service as eo
    p = _bank_payload()
    res = eo.recompute(p, "D05.SI", {"shared": {"bank_nim": 0.0225, "bank_credit_cost_bps": 30.0}})
    b = res["scenarios"]["base"]
    assert res["model_kind"] == "bank" and b["three_statements"]["kind"] == "bank" and b["three_statements"]["overrides_applied"] == {"nim_fy1": 0.0225, "nim_fy2": 0.0225, "credit_cost_bps_fy1": 30.0, "credit_cost_bps_fy2": 30.0}
    assert b["three_statements"]["assumptions"]["nim_fy1"]["source"] == "user override"
    legs = b["legs"]
    assert set(legs) == {"GGM (P/B)", "Forward P/E"}
    g = legs["GGM (P/B)"]
    assert 0.3 <= g["multiple"] <= 4.0 and g["metric_after"] == pytest.approx(b["three_statements"]["bvps_fy1"]) and g["value_after"] == pytest.approx(g["metric_after"] * g["multiple"])
    assert legs["Forward P/E"]["metric_after"] == pytest.approx(b["three_statements"]["eps_fy1"]) and legs["Forward P/E"]["value_after"] == pytest.approx(b["three_statements"]["eps_fy1"] * 10.2)
    assert b["intrinsic_value"] == pytest.approx(0.45 * g["value_after"] + 0.25 * 36.8 + 0.15 * legs["Forward P/E"]["value_after"] + 0.15 * 46.8)
    assert b["target"] == pytest.approx(77.2 + 0.5 * (b["intrinsic_value"] - 77.2)) and res["after"]["12m_price_target"] is not None
    out = eo.apply_to_payload(copy.deepcopy(p), "D05.SI", {"id": "o", "created_at": "2026-10-03T00:00:00", "note": "wider NIM", "overrides": res["overrides"], "result": res})
    dr_ = out["data"]["dcf_range"]["D05.SI"]
    assert dr_["base"]["leg_inputs"]["GGM (P/B)"]["user_override"] is True and dr_["base"]["leg_inputs"]["GGM (P/B)"]["assumptions"]["bvps_basis"].startswith("user estimate")
    assert dr_["base"]["method_iv_table"]["GGM (P/B)"] == pytest.approx(g["value_after"], abs=0.01) and dr_["three_statements"]["override"]["note"] == "wider NIM"
    assert dr_["three_statements_scenarios"]["bear"]["kind"] == "bank"
    with pytest.raises(ValueError, match="outside"):
        eo.normalize_overrides({"shared": {"bank_nim": 0.5}})
    res0 = eo.recompute(p, "D05.SI", {"shared": {"wacc": 0.10}})                 # no bank driver changed: nothing re-prices
    assert res0["scenarios"]["base"]["legs"] == {} and "no driver changed" in res0["scenarios"]["base"]["note"]


def test_the_model_tab_and_the_pdf_print_the_bank_layout():
    from src.utils.valuation_workbook import build_workbook
    p = _bank_payload()
    wb = load_workbook(io.BytesIO(build_workbook(p, "D05.SI", load_statements=_wbt._statements)))
    ws = wb["Model"]
    col_a = [str(c.value) for c in ws["A"] if c.value is not None]
    for head in ("Assumptions (each with its source)", "Balance sheet drivers", "Income statement", "Capital", "Per share and returns", "Reconciliation suite (engine build)"):
        assert head in col_a
    assert "Net interest margin" in col_a and "CET1 ratio" in col_a and "RoTE" in col_a and str(ws["D4"].value) == "FY2026E"
    assert any(v.startswith("Suite result: ALL OK") for v in col_a) and any("vs cost of equity" in v for v in col_a)
    # the IS tab's forecast columns carry the bank model's lines; the BS tab its total assets and equity; the CFS tab has none
    is_ws = wb["IS"]
    rr = [c.row for c in is_ws["A"] if c.value == "Net interest income"]
    f0 = next(c for c in range(3, 20) if str(is_ws.cell(row=4, column=c).value).startswith("=EDATE("))
    assert rr and str(is_ws.cell(row=rr[-1], column=f0).value).startswith("='Model'!")
    rev_r = [c.row for c in is_ws["A"] if c.value == "Revenue"][0]
    assert str(is_ws.cell(row=rev_r, column=f0).value).startswith("='Model'!")                     # total income on the revenue row
    bs_ws = wb["BS"]
    rr = [c.row for c in bs_ws["A"] if c.value == "Shareholders' equity"]
    f0b = next(c for c in range(3, 20) if str(bs_ws.cell(row=4, column=c).value).startswith("=EDATE("))
    assert rr and str(bs_ws.cell(row=rr[-1], column=f0b).value).startswith("='Model'!")
    assert wb["CFS"].cell(row=4, column=8).value is None
    from src.utils import pdf_report as pr
    from reportlab.lib.styles import getSampleStyleSheet
    flow = pr._three_statement_block_pdf(p["data"]["dcf_range"]["D05.SI"], getSampleStyleSheet(), 500.0)
    texts = " ".join(getattr(f, "text", "") for f in flow)
    assert "Bank earnings-and-capital model" in texts and "guided: loan_growth_fy1" in texts
    assert len([f for f in flow if f.__class__.__name__ == "Table"]) == 2


@pytest.mark.slow
def test_the_dbs_golden_keeps_the_research_roe_without_family_guidance_and_reads_the_model_with_it():
    """Without management's drivers in the research the model is a cross-check (02888.HK fell 17% when a
    line-item NIM re-priced its GGM); with family guidance the legs read the model's RoTE and FY+1 book."""
    from src.memory.golden_replay import replay_fixture
    out = replay_fixture("D05_SI")
    dr_ = out["entry"]
    th = dr_.get("three_statements") or {}
    assert th.get("kind") == "bank" and th.get("fy_labels") and th["reconciliation"]["ok"] and th.get("feeds_legs") is False
    ggm = dr_["base"]["leg_inputs"]["GGM (P/B)"]
    assert ggm["assumptions"]["bvps_basis"] == "latest book" and not any("bank model" in s for s in ggm["assumptions"]["provenance"])
    assert any(f.startswith("Bank earnings-and-capital model") and "cross-check" in f for f in dr_["base"]["forward_flags"])
    assert out["base_iv"] == pytest.approx(44.44, abs=0.05)                     # unchanged: the golden holds
    block = {"fiscal_year_1": "FY2026", "fiscal_year_2": "FY2027", "confidence": "HIGH", "estimates": {"base": {"revenue_growth_fy1": 0.022, "revenue_growth_fy2": 0.035}},
             "family_metrics": {"family": "bank", "bank": {"loan_growth_fy1": 0.05, "nim_fy1": 0.0205, "cost_to_income_fy1": 0.40, "credit_cost_bps_fy1": 20.0, "payout_ratio": 0.55}}}
    out2 = replay_fixture("D05_SI", state_patch={"guidance_estimates": {"D05.SI": block}})
    dr2 = out2["entry"]
    th2 = dr2["three_statements"]
    assert th2["feeds_legs"] is True and set(th2["guided_fields"]) == {"loan_growth_fy1", "nim_fy1", "cost_to_income_fy1", "credit_cost_bps_fy1", "payout_ratio"}
    ggm2 = dr2["base"]["leg_inputs"]["GGM (P/B)"]
    assert ggm2["assumptions"]["bvps_basis"].startswith("bank model") and ggm2["assumptions"]["bvps"] == pytest.approx(th2["bvps_fy1"])
    assert any("bank model" in s for s in ggm2["assumptions"]["provenance"])
    assert dr2.get("guidance_forecast") in (None, {})                           # the working-capital roll does not run for a bank
    fpe = dr2["base"]["leg_inputs"].get("Forward P/E") or {}
    assert fpe.get("metric_value") == pytest.approx(th2["eps_fy1"]) if fpe else True
