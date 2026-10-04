"""Workbook tie-out: rebuild a run's valuation workbook, recalculate every formula, and list
each Check cell against the engine (plan 2026-10, Phase 0 step 5).

    python scripts/workbook_tieout.py RUN.json [RUN2.json ...] [--statements] [--keep DIR]

RUN.json is {"run_id", "run_at", "data": {...}} -- the web_runs payload. Without --statements
the IS/BS/CFS tabs are left empty (no FMP calls); the valuation tabs do not depend on them.
A Check passes when |value| <= 0.01 (one published rounding unit). Exit status 1 when any fails.
"""
from __future__ import annotations

import argparse
import io
import json
import os
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

TOL = 0.01


def _checks(xlsx_path: str) -> list[tuple[str, str, object]]:
    """[(sheet, cell, value)] for every recalculated cell labelled Check or reading OK/REVIEW."""
    import formulas
    from openpyxl import load_workbook

    wb = load_workbook(xlsx_path)
    targets: list[tuple[str, str]] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if not isinstance(v, str):
                    continue
                if v.strip().startswith("Check"):
                    # the value sits to the right of the label (same row), first formula cell
                    for cc in row[c.column:]:
                        if isinstance(cc.value, str) and cc.value.startswith("="):
                            targets.append((ws.title, cc.coordinate))
                            break
                elif (v.startswith("=IF(ABS(") or v.startswith("=IF(ROUND(ABS(")) and '"REVIEW"' in v:
                    targets.append((ws.title, c.coordinate))
                elif ws.title == "Multiples" and c.column == 17 and v.startswith("="):
                    targets.append((ws.title, c.coordinate))
    model = formulas.ExcelModel().loads(xlsx_path).finish()
    sol = {k.upper(): v for k, v in model.calculate().items()}
    name = Path(xlsx_path).name.upper()
    out = []
    for sheet, cell in targets:
        key = f"'[{name}]{sheet.upper()}'!{cell}".upper()
        val = sol.get(key)
        v = getattr(val, "value", val)
        try:
            v = v[0][0] if hasattr(v, "__getitem__") and not isinstance(v, str) else v
        except Exception:  # noqa: BLE001
            pass
        out.append((sheet, cell, v))
    return out


def tieout(run_path: str, statements: bool = False, keep: str | None = None) -> int:
    from src.utils.valuation_workbook import build_workbook

    run = json.load(open(run_path, encoding="utf-8"))
    data = run.get("data") or {}
    fails = 0
    for ticker in (data.get("dcf_range") or {}):
        loaders = {}
        if statements:
            from app.backend.services.workbook_export_service import _members, load_statements
            loaders = {"load_members": _members, "load_statements": load_statements}
        blob = build_workbook(run, ticker, **loaders)
        d = keep or tempfile.mkdtemp()
        path = os.path.join(d, f"{ticker.replace('.', '_')}_tieout.xlsx")
        Path(path).write_bytes(blob)
        rows = _checks(path)
        bad = []
        for sheet, cell, v in rows:
            ok = (isinstance(v, str) and v.strip() in ("OK", "ALL OK")) or (
                isinstance(v, (int, float)) and abs(float(v)) <= TOL)
            if not ok:
                bad.append((sheet, cell, v))
        fails += len(bad)
        print(f"{ticker}: {len(rows) - len(bad)}/{len(rows)} checks pass  ({path})")
        for sheet, cell, v in bad:
            print(f"   FAIL {sheet}!{cell} = {v}")
    return fails


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("runs", nargs="+")
    ap.add_argument("--statements", action="store_true")
    ap.add_argument("--keep")
    a = ap.parse_args()
    total = sum(tieout(p, a.statements, a.keep) for p in a.runs)
    sys.exit(1 if total else 0)


if __name__ == "__main__":
    main()
